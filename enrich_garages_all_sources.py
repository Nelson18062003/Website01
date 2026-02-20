#!/usr/bin/env python3
"""
Enrich garages_auto_zones_france.xlsx with phone numbers from multiple sources.

Strategy (by priority):
1. OSM SIRET exact match  → highest confidence, zero ambiguity
2. Pages Jaunes name+CP   → high confidence (same postal code + similar name)
3. OSM name+CP match      → medium confidence (fuzzy matching)

Sources:
- osm_garages_phones.json: 12,280 garages with phone, 4,358 with SIRET
- pagesjaunes_garages_cache/: Scraped Pages Jaunes data per department

Output: Updated garages_auto_zones_france.xlsx with phone numbers filled in.
"""

import json
import os
import re
import sys
import openpyxl
from openpyxl.styles import Font, Alignment

# =============================================================================
# Configuration
# =============================================================================

EXCEL_FILE = "garages_auto_zones_france.xlsx"
OSM_FILE = "osm_garages_phones.json"
PJ_CACHE_DIR = "pagesjaunes_garages_cache"
OUTPUT_FILE = "garages_auto_zones_france.xlsx"

# Column indices in the Excel file (1-based)
COL_NUM = 1
COL_NAME = 2
COL_ADDR = 3
COL_CP = 4
COL_CITY = 5
COL_DEPT = 6
COL_PHONE_INTL = 10
COL_PHONE_ORIG = 11
COL_EMAIL = 12
COL_SIRET = 14


# =============================================================================
# Name normalization
# =============================================================================

def normalize_name(name):
    """Normalize a business name for comparison."""
    if not name:
        return ""
    n = str(name).lower().strip()
    # Remove parenthesized content
    n = re.sub(r'\([^)]*\)', '', n)
    # Remove special chars
    n = re.sub(r'[^a-z0-9àâäéèêëïîôùûüÿçœæ\s]', ' ', n)
    # Normalize accents
    for src, dst in [('é', 'e'), ('è', 'e'), ('ê', 'e'), ('ë', 'e'),
                     ('à', 'a'), ('â', 'a'), ('ä', 'a'),
                     ('ù', 'u'), ('û', 'u'), ('ü', 'u'),
                     ('î', 'i'), ('ï', 'i'), ('ô', 'o'), ('ç', 'c'),
                     ('œ', 'oe'), ('æ', 'ae'), ('ÿ', 'y')]:
        n = n.replace(src, dst)
    n = re.sub(r'\s+', ' ', n).strip()
    return n


STOP_WORDS = {
    "garage", "garages", "auto", "automobile", "automobiles", "sarl", "sas",
    "eurl", "sa", "srl", "et", "de", "du", "des", "le", "la", "les", "l",
    "un", "une", "en", "au", "aux", "par", "pour", "sur", "chez", "fils",
    "ets", "etablissements", "monsieur", "madame", "mr", "mme", "m",
    "reparation", "entretien", "mecanique", "carrosserie", "service",
    "services", "centre", "atelier", "station", "point", "agent",
    "concessionnaire", "concession", "vente", "france", "international",
    "group", "groupe", "societe", "ste",
}


def name_tokens(name):
    """Get significant tokens from a normalized name."""
    normalized = normalize_name(name)
    tokens = set(normalized.split())
    return {t for t in tokens if t not in STOP_WORDS and len(t) > 1}


def extract_commercial_name(name):
    """Extract commercial name from parentheses if present.
    e.g. 'BENELO AUTO (MIDAS)' → 'MIDAS'
    e.g. 'SAS HURELLE MEYLAN (HURELLE MEYLAN)' → 'HURELLE MEYLAN'
    """
    if not name:
        return ""
    match = re.search(r'\(([^)]+)\)', str(name))
    return match.group(1).strip() if match else ""


def name_similarity(name1, name2):
    """Token-based Jaccard similarity, also checking commercial names in parens.
    For enterprise names like 'BENELO AUTO (MIDAS)', also tries matching
    the commercial name 'MIDAS' against the other name.
    Returns the best score found.
    """
    tokens1 = name_tokens(name1)
    tokens2 = name_tokens(name2)

    scores = []

    # Standard comparison
    if tokens1 and tokens2:
        intersection = tokens1 & tokens2
        union = tokens1 | tokens2
        scores.append(len(intersection) / len(union) if union else 0.0)

    # Try commercial name from name1 (enterprise side)
    commercial1 = extract_commercial_name(name1)
    if commercial1:
        tc1 = name_tokens(commercial1)
        if tc1 and tokens2:
            intersection = tc1 & tokens2
            union = tc1 | tokens2
            scores.append(len(intersection) / len(union) if union else 0.0)

    # Try commercial name from name2
    commercial2 = extract_commercial_name(name2)
    if commercial2:
        tc2 = name_tokens(commercial2)
        if tokens1 and tc2:
            intersection = tokens1 & tc2
            union = tokens1 | tc2
            scores.append(len(intersection) / len(union) if union else 0.0)

    return max(scores) if scores else 0.0


def format_phone_international(phone):
    """Convert French phone to international +33 format."""
    if not phone:
        return ""
    p = str(phone).strip()
    for sep in [";", "/", ","]:
        if sep in p:
            p = p.split(sep)[0].strip()
    digits = re.sub(r'[^\d+]', '', p)
    if digits.startswith('+33'):
        digits = '0' + digits[3:]
    elif digits.startswith('0033'):
        digits = '0' + digits[4:]
    elif digits.startswith('33') and len(digits) == 11:
        digits = '0' + digits[2:]
    digits = re.sub(r'[^\d]', '', digits)
    if digits.startswith('0') and len(digits) == 10:
        return f"+33 {digits[1]} {digits[2:4]} {digits[4:6]} {digits[6:8]} {digits[8:10]}"
    elif len(digits) == 9 and not digits.startswith('0'):
        return f"+33 {digits[0]} {digits[1:3]} {digits[3:5]} {digits[5:7]} {digits[7:9]}"
    return p


def get_dept_from_postal(postal_code):
    """Extract department number from postal code."""
    cp = str(postal_code or "").strip()
    if not cp:
        return ""
    if cp[:3] in ("971", "972", "973", "974", "976"):
        return cp[:3]
    if len(cp) >= 2:
        return cp[:2]
    return ""


# =============================================================================
# Data loading
# =============================================================================

def load_excel_garages(filepath):
    """Load garages from the Excel file."""
    print(f"  Chargement de {filepath}...")
    wb = openpyxl.load_workbook(filepath)
    ws = wb.active

    garages = []
    for row in range(2, ws.max_row + 1):
        siret = str(ws.cell(row=row, column=COL_SIRET).value or "").strip()
        phone = str(ws.cell(row=row, column=COL_PHONE_INTL).value or "").strip()
        garages.append({
            "row": row,
            "name": str(ws.cell(row=row, column=COL_NAME).value or ""),
            "address": str(ws.cell(row=row, column=COL_ADDR).value or ""),
            "postal_code": str(ws.cell(row=row, column=COL_CP).value or ""),
            "city": str(ws.cell(row=row, column=COL_CITY).value or ""),
            "dept": str(ws.cell(row=row, column=COL_DEPT).value or ""),
            "siret": siret,
            "has_phone": phone not in ("", "None"),
        })

    print(f"    {len(garages):,} garages chargés")
    already_phone = sum(1 for g in garages if g["has_phone"])
    print(f"    Déjà avec téléphone: {already_phone:,}")
    return garages, wb, ws


def load_osm_data(filepath):
    """Load OSM garage data and index by SIRET and postal code."""
    print(f"  Chargement de {filepath}...")
    with open(filepath, "r", encoding="utf-8") as f:
        elements = json.load(f)

    by_siret = {}
    by_postcode = {}
    total_with_phone = 0
    total_with_siret = 0

    for elem in elements:
        tags = elem.get("tags", {})
        name = tags.get("name", "")
        phone = tags.get("phone", "") or tags.get("contact:phone", "")
        email = tags.get("email", "") or tags.get("contact:email", "")
        siret = tags.get("ref:FR:SIRET", "") or tags.get("siret", "")
        postcode = tags.get("addr:postcode", "")
        city = tags.get("addr:city", "") or tags.get("addr:commune", "")
        street = tags.get("addr:street", "")
        website = tags.get("website", "") or tags.get("contact:website", "")

        if not phone:
            continue

        total_with_phone += 1

        entry = {
            "name": name,
            "phone": phone,
            "email": email,
            "siret": siret.replace(" ", ""),
            "postcode": postcode,
            "city": city,
            "street": street,
            "website": website,
            "osm_id": elem.get("id", ""),
        }

        if siret:
            clean_siret = siret.replace(" ", "")
            by_siret[clean_siret] = entry
            total_with_siret += 1

        if postcode:
            by_postcode.setdefault(postcode, []).append(entry)

    print(f"    {total_with_phone:,} entrées avec téléphone")
    print(f"    {total_with_siret:,} entrées avec SIRET")
    print(f"    {len(by_postcode):,} codes postaux uniques")
    return by_siret, by_postcode


def normalize_city(city):
    """Normalize city name for matching."""
    if not city:
        return ""
    c = normalize_name(city)
    c = re.sub(r'\bsaint\b', 'st', c)
    c = re.sub(r'\bsainte\b', 'ste', c)
    return c


def load_pj_data(cache_dir):
    """Load Pages Jaunes scraped data from cache directory."""
    print(f"  Chargement des données Pages Jaunes ({cache_dir})...")
    by_postcode = {}
    by_city = {}
    total = 0
    total_with_phone = 0

    if not os.path.isdir(cache_dir):
        print(f"    [WARN] Répertoire {cache_dir} non trouvé.")
        return by_postcode, by_city

    for filename in sorted(os.listdir(cache_dir)):
        if not filename.endswith(".json"):
            continue
        filepath = os.path.join(cache_dir, filename)
        with open(filepath, "r", encoding="utf-8") as f:
            entries = json.load(f)
        for entry in entries:
            total += 1
            if entry.get("phone"):
                total_with_phone += 1
            cp = entry.get("postal_code", "")
            city = entry.get("city", "")
            if cp:
                by_postcode.setdefault(cp, []).append(entry)
            if city:
                city_norm = normalize_city(city)
                if city_norm:
                    by_city.setdefault(city_norm, []).append(entry)

    print(f"    {total:,} entrées chargées")
    print(f"    {total_with_phone:,} avec téléphone")
    print(f"    {len(by_postcode):,} codes postaux uniques")
    print(f"    {len(by_city):,} villes uniques")
    return by_postcode, by_city


# =============================================================================
# Matching engine
# =============================================================================

def find_best_pj_match(garage_name, garage_addr, candidates, threshold=0.40):
    """Find the best matching PJ entry by name similarity."""
    best_score = 0
    best_entry = None

    for entry in candidates:
        if not entry.get("phone"):
            continue

        pj_name = entry.get("name", "")
        score = name_similarity(garage_name, pj_name)

        # Bonus for address similarity
        pj_addr = entry.get("address", "") or entry.get("full_address", "")
        if garage_addr and pj_addr:
            addr_tokens_g = set(normalize_name(garage_addr).split())
            addr_tokens_p = set(normalize_name(pj_addr).split())
            common = addr_tokens_g & addr_tokens_p
            if len(common) >= 2:
                score += 0.15

        if score > best_score:
            best_score = score
            best_entry = entry

    if best_score >= threshold and best_entry:
        return best_entry, best_score
    return None, 0.0


def find_best_osm_match(garage_name, candidates, threshold=0.40):
    """Find the best matching OSM entry by name similarity."""
    best_score = 0
    best_entry = None

    for entry in candidates:
        if not entry.get("phone"):
            continue
        score = name_similarity(garage_name, entry.get("name", ""))
        if score > best_score:
            best_score = score
            best_entry = entry

    if best_score >= threshold and best_entry:
        return best_entry, best_score
    return None, 0.0


# =============================================================================
# Main enrichment
# =============================================================================

def enrich(garages, osm_by_siret, osm_by_postcode, pj_by_postcode, pj_by_city, wb, ws):
    """Run multi-source enrichment on garages."""

    stats = {
        "already_had": 0,
        "osm_siret": 0,
        "pj_name_cp": 0,
        "pj_name_city": 0,
        "osm_name_cp": 0,
        "not_found": 0,
    }

    total = len(garages)
    used_osm_ids = set()
    used_pj_ids = set()

    for i, garage in enumerate(garages):
        if garage["has_phone"]:
            stats["already_had"] += 1
            continue

        phone = ""
        phone_intl = ""
        email = ""
        source_tag = ""

        # ----- Strategy 1: OSM SIRET exact match -----
        siret = garage["siret"]
        if siret and siret in osm_by_siret:
            osm_entry = osm_by_siret[siret]
            osm_id = osm_entry.get("osm_id", "")
            if osm_id not in used_osm_ids:
                phone = osm_entry["phone"]
                phone_intl = format_phone_international(phone)
                email = osm_entry.get("email", "")
                source_tag = "OSM (SIRET exact)"
                used_osm_ids.add(osm_id)
                stats["osm_siret"] += 1

        # ----- Strategy 2a: Pages Jaunes name+CP match -----
        if not phone:
            cp = garage["postal_code"]
            if cp and cp in pj_by_postcode:
                available = [e for e in pj_by_postcode[cp]
                             if e.get("pj_id", "") not in used_pj_ids]
                match, score = find_best_pj_match(
                    garage["name"], garage["address"], available, threshold=0.40
                )
                if match:
                    phone = match["phone"]
                    phone_intl = match.get("phone_intl", "") or format_phone_international(phone)
                    email = match.get("email", "") or ""
                    pj_id = match.get("pj_id", "")
                    if pj_id:
                        used_pj_ids.add(pj_id)
                    source_tag = f"PJ (nom+CP, {score:.2f})"
                    stats["pj_name_cp"] += 1

        # ----- Strategy 2b: Pages Jaunes name+city match -----
        # Same city may have multiple postal codes, so match by city name
        if not phone:
            city = garage["city"]
            if city:
                city_norm = normalize_city(city)
                if city_norm and city_norm in pj_by_city:
                    available = [e for e in pj_by_city[city_norm]
                                 if e.get("pj_id", "") not in used_pj_ids]
                    match, score = find_best_pj_match(
                        garage["name"], garage["address"], available, threshold=0.45
                    )
                    if match:
                        phone = match["phone"]
                        phone_intl = match.get("phone_intl", "") or format_phone_international(phone)
                        email = match.get("email", "") or ""
                        pj_id = match.get("pj_id", "")
                        if pj_id:
                            used_pj_ids.add(pj_id)
                        source_tag = f"PJ (nom+ville, {score:.2f})"
                        stats["pj_name_city"] += 1

        # ----- Strategy 3: OSM name+CP match -----
        if not phone:
            cp = garage["postal_code"]
            if cp and cp in osm_by_postcode:
                available = [e for e in osm_by_postcode[cp]
                             if e.get("osm_id", "") not in used_osm_ids]
                match, score = find_best_osm_match(
                    garage["name"], available, threshold=0.40
                )
                if match:
                    phone = match["phone"]
                    phone_intl = format_phone_international(phone)
                    email = match.get("email", "")
                    osm_id = match.get("osm_id", "")
                    if osm_id:
                        used_osm_ids.add(osm_id)
                    source_tag = f"OSM (nom+CP, {score:.2f})"
                    stats["osm_name_cp"] += 1

        # ----- Update Excel -----
        if phone:
            row = garage["row"]
            ws.cell(row=row, column=COL_PHONE_INTL, value=phone_intl)
            ws.cell(row=row, column=COL_PHONE_ORIG, value=phone)
            if email:
                current_email = str(ws.cell(row=row, column=COL_EMAIL).value or "")
                if not current_email or current_email == "None":
                    ws.cell(row=row, column=COL_EMAIL, value=email)
        else:
            stats["not_found"] += 1

        # Progress
        if (i + 1) % 5000 == 0 or (i + 1) == total:
            found = stats["osm_siret"] + stats["pj_name_cp"] + stats["pj_name_city"] + stats["osm_name_cp"]
            pct = found / total * 100
            print(f"    [{i+1:,}/{total:,}] Enrichis: {found:,} ({pct:.1f}%)")

    return stats


def main():
    print("=" * 70)
    print("  ENRICHISSEMENT GARAGES - MULTI-SOURCES")
    print("  Sources: OSM (SIRET exact) + Pages Jaunes + OSM (nom+CP)")
    print("=" * 70)
    print()

    # Load data
    print("PHASE 1: Chargement des données")
    print("-" * 50)
    garages, wb, ws = load_excel_garages(EXCEL_FILE)
    osm_by_siret, osm_by_postcode = load_osm_data(OSM_FILE)
    pj_by_postcode, pj_by_city = load_pj_data(PJ_CACHE_DIR)

    # Enrich
    print(f"\nPHASE 2: Enrichissement")
    print("-" * 50)
    stats = enrich(garages, osm_by_siret, osm_by_postcode, pj_by_postcode, pj_by_city, wb, ws)

    # Update summary sheet
    if "Résumé" in wb.sheetnames:
        ws_s = wb["Résumé"]
        total_phones = (stats["already_had"] + stats["osm_siret"]
                        + stats["pj_name_cp"] + stats["pj_name_city"]
                        + stats["osm_name_cp"])
        for row in range(1, ws_s.max_row + 1):
            val = ws_s.cell(row=row, column=1).value
            if val and "téléphone" in str(val).lower():
                ws_s.cell(row=row, column=2, value=total_phones)
                break

    # Save
    print(f"\nPHASE 3: Sauvegarde")
    print("-" * 50)
    wb.save(OUTPUT_FILE)
    print(f"  Fichier sauvegardé: {OUTPUT_FILE}")

    # Summary
    total_found = (stats["osm_siret"] + stats["pj_name_cp"]
                   + stats["pj_name_city"] + stats["osm_name_cp"])
    grand_total = total_found + stats["already_had"]
    total_garages = len(garages)

    print(f"\n{'=' * 70}")
    print(f"  RÉSUMÉ FINAL")
    print(f"{'=' * 70}")
    print(f"  Total garages:              {total_garages:,}")
    print(f"  Déjà avec téléphone:        {stats['already_had']:,}")
    print(f"  Nouveaux téléphones:         {total_found:,}")
    print(f"    OSM (SIRET exact):         {stats['osm_siret']:,}")
    print(f"    Pages Jaunes (nom+CP):     {stats['pj_name_cp']:,}")
    print(f"    Pages Jaunes (nom+ville):  {stats['pj_name_city']:,}")
    print(f"    OSM (nom+CP):              {stats['osm_name_cp']:,}")
    print(f"  Sans téléphone:             {stats['not_found']:,}")
    print(f"  -------")
    print(f"  Couverture totale:          {grand_total:,} / {total_garages:,} "
          f"({grand_total/total_garages*100:.1f}%)")
    print(f"\n  Fichier: {OUTPUT_FILE}")
    print("\nTerminé!")


if __name__ == "__main__":
    main()
