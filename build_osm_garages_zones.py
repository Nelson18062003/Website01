#!/usr/bin/env python3
"""
Build an Excel list of OSM garages with phone numbers,
filtered by Top Zone and Middle Zone departments.

Steps:
1. Load OSM garage data (with phone numbers)
2. For entries with postcode → extract department
3. For entries without postcode but with lat/lon → batch reverse geocode
4. Filter to Top Zone + Middle Zone departments only
5. Write formatted Excel file
"""

import json
import re
import csv
import io
import time
import requests
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

# =============================================================================
# Configuration
# =============================================================================

OSM_FILE = "osm_garages_phones.json"
OUTPUT_FILE = "garages_osm_top_middle_zones.xlsx"

GEOCODE_BATCH_URL = "https://api-adresse.data.gouv.fr/reverse/csv/"
BATCH_SIZE = 1000  # Keep batches manageable

# Department definitions
DEPARTMENTS = {
    "01": ("Ain", "Auvergne-Rhône-Alpes"),
    "02": ("Aisne", "Hauts-de-France"),
    "03": ("Allier", "Auvergne-Rhône-Alpes"),
    "04": ("Alpes-de-Haute-Provence", "Provence-Alpes-Côte d'Azur"),
    "05": ("Hautes-Alpes", "Provence-Alpes-Côte d'Azur"),
    "06": ("Alpes-Maritimes", "Provence-Alpes-Côte d'Azur"),
    "07": ("Ardèche", "Auvergne-Rhône-Alpes"),
    "08": ("Ardennes", "Grand Est"),
    "09": ("Ariège", "Occitanie"),
    "10": ("Aube", "Grand Est"),
    "11": ("Aude", "Occitanie"),
    "12": ("Aveyron", "Occitanie"),
    "13": ("Bouches-du-Rhône", "Provence-Alpes-Côte d'Azur"),
    "14": ("Calvados", "Normandie"),
    "15": ("Cantal", "Auvergne-Rhône-Alpes"),
    "16": ("Charente", "Nouvelle-Aquitaine"),
    "17": ("Charente-Maritime", "Nouvelle-Aquitaine"),
    "18": ("Cher", "Centre-Val de Loire"),
    "19": ("Corrèze", "Nouvelle-Aquitaine"),
    "2A": ("Corse-du-Sud", "Corse"),
    "2B": ("Haute-Corse", "Corse"),
    "21": ("Côte-d'Or", "Bourgogne-Franche-Comté"),
    "22": ("Côtes-d'Armor", "Bretagne"),
    "23": ("Creuse", "Nouvelle-Aquitaine"),
    "24": ("Dordogne", "Nouvelle-Aquitaine"),
    "25": ("Doubs", "Bourgogne-Franche-Comté"),
    "26": ("Drôme", "Auvergne-Rhône-Alpes"),
    "27": ("Eure", "Normandie"),
    "28": ("Eure-et-Loir", "Centre-Val de Loire"),
    "29": ("Finistère", "Bretagne"),
    "30": ("Gard", "Occitanie"),
    "31": ("Haute-Garonne", "Occitanie"),
    "32": ("Gers", "Occitanie"),
    "33": ("Gironde", "Nouvelle-Aquitaine"),
    "34": ("Hérault", "Occitanie"),
    "35": ("Ille-et-Vilaine", "Bretagne"),
    "36": ("Indre", "Centre-Val de Loire"),
    "37": ("Indre-et-Loire", "Centre-Val de Loire"),
    "38": ("Isère", "Auvergne-Rhône-Alpes"),
    "39": ("Jura", "Bourgogne-Franche-Comté"),
    "40": ("Landes", "Nouvelle-Aquitaine"),
    "41": ("Loir-et-Cher", "Centre-Val de Loire"),
    "42": ("Loire", "Auvergne-Rhône-Alpes"),
    "43": ("Haute-Loire", "Auvergne-Rhône-Alpes"),
    "44": ("Loire-Atlantique", "Pays de la Loire"),
    "45": ("Loiret", "Centre-Val de Loire"),
    "46": ("Lot", "Occitanie"),
    "47": ("Lot-et-Garonne", "Nouvelle-Aquitaine"),
    "48": ("Lozère", "Occitanie"),
    "49": ("Maine-et-Loire", "Pays de la Loire"),
    "50": ("Manche", "Normandie"),
    "51": ("Marne", "Grand Est"),
    "52": ("Haute-Marne", "Grand Est"),
    "53": ("Mayenne", "Pays de la Loire"),
    "54": ("Meurthe-et-Moselle", "Grand Est"),
    "55": ("Meuse", "Grand Est"),
    "56": ("Morbihan", "Bretagne"),
    "57": ("Moselle", "Grand Est"),
    "58": ("Nièvre", "Bourgogne-Franche-Comté"),
    "59": ("Nord", "Hauts-de-France"),
    "60": ("Oise", "Hauts-de-France"),
    "61": ("Orne", "Normandie"),
    "62": ("Pas-de-Calais", "Hauts-de-France"),
    "63": ("Puy-de-Dôme", "Auvergne-Rhône-Alpes"),
    "64": ("Pyrénées-Atlantiques", "Nouvelle-Aquitaine"),
    "65": ("Hautes-Pyrénées", "Occitanie"),
    "66": ("Pyrénées-Orientales", "Occitanie"),
    "67": ("Bas-Rhin", "Grand Est"),
    "68": ("Haut-Rhin", "Grand Est"),
    "69": ("Rhône", "Auvergne-Rhône-Alpes"),
    "70": ("Haute-Saône", "Bourgogne-Franche-Comté"),
    "71": ("Saône-et-Loire", "Bourgogne-Franche-Comté"),
    "72": ("Sarthe", "Pays de la Loire"),
    "73": ("Savoie", "Auvergne-Rhône-Alpes"),
    "74": ("Haute-Savoie", "Auvergne-Rhône-Alpes"),
    "75": ("Paris", "Île-de-France"),
    "76": ("Seine-Maritime", "Normandie"),
    "77": ("Seine-et-Marne", "Île-de-France"),
    "78": ("Yvelines", "Île-de-France"),
    "79": ("Deux-Sèvres", "Nouvelle-Aquitaine"),
    "80": ("Somme", "Hauts-de-France"),
    "81": ("Tarn", "Occitanie"),
    "82": ("Tarn-et-Garonne", "Occitanie"),
    "83": ("Var", "Provence-Alpes-Côte d'Azur"),
    "84": ("Vaucluse", "Provence-Alpes-Côte d'Azur"),
    "85": ("Vendée", "Pays de la Loire"),
    "86": ("Vienne", "Nouvelle-Aquitaine"),
    "87": ("Haute-Vienne", "Nouvelle-Aquitaine"),
    "88": ("Vosges", "Grand Est"),
    "89": ("Yonne", "Bourgogne-Franche-Comté"),
    "90": ("Territoire de Belfort", "Bourgogne-Franche-Comté"),
    "91": ("Essonne", "Île-de-France"),
    "92": ("Hauts-de-Seine", "Île-de-France"),
    "93": ("Seine-Saint-Denis", "Île-de-France"),
    "94": ("Val-de-Marne", "Île-de-France"),
    "95": ("Val-d'Oise", "Île-de-France"),
    "971": ("Guadeloupe", "Guadeloupe"),
    "972": ("Martinique", "Martinique"),
    "973": ("Guyane", "Guyane"),
    "974": ("La Réunion", "La Réunion"),
    "976": ("Mayotte", "Mayotte"),
}

TOP_ZONE_DEPTS = {"06", "59", "60", "77", "78", "88", "91", "93", "95"}
MIDDLE_ZONE_DEPTS = {
    "17", "24", "27", "28", "30", "31", "33", "35", "38", "40", "41", "42",
    "44", "45", "46", "57", "62", "67", "72", "73", "76", "79", "80", "83",
    "86", "971", "973",
}
ALL_TARGET_DEPTS = TOP_ZONE_DEPTS | MIDDLE_ZONE_DEPTS


# =============================================================================
# Helpers
# =============================================================================

def get_dept_from_postcode(pc):
    """Extract department number from postal code."""
    pc = str(pc or "").strip()
    if not pc:
        return ""
    if pc[:3] in ("971", "972", "973", "974", "976"):
        return pc[:3]
    if len(pc) >= 2:
        return pc[:2]
    return ""


def get_zone(dept):
    """Get zone label."""
    if dept in TOP_ZONE_DEPTS:
        return "Top Zone"
    elif dept in MIDDLE_ZONE_DEPTS:
        return "Middle Zone"
    return ""


def format_phone_international(phone):
    """Convert French phone to +33 X XX XX XX XX format."""
    if not phone:
        return ""
    p = str(phone).strip()
    if ';' in p:
        p = p.split(';')[0].strip()
    if '/' in p:
        p = p.split('/')[0].strip()

    digits = re.sub(r'[^\d+]', '', p)
    if digits.startswith('+33'):
        digits = '0' + digits[3:]
    elif digits.startswith('0033'):
        digits = '0' + digits[4:]
    elif digits.startswith('33') and len(digits) == 11:
        digits = '0' + digits[2:]

    digits = re.sub(r'[^\d]', '', digits)

    if digits.startswith('0') and len(digits) == 10:
        return f"+33 {digits[1]} {digits[2:4]} {digits[4:6]} {digits[6:8]} {digits[8:10]}"
    elif len(digits) == 9 and not digits.startswith('0'):
        return f"+33 {digits[0]} {digits[1:3]} {digits[3:5]} {digits[5:7]} {digits[7:9]}"

    return p


# =============================================================================
# Reverse geocoding
# =============================================================================

def batch_reverse_geocode(entries_with_coords):
    """Batch reverse geocode entries using api-adresse.data.gouv.fr."""
    if not entries_with_coords:
        return {}

    results = {}  # osm_id → {"postcode": ..., "city": ..., "dept": ...}
    total = len(entries_with_coords)
    processed = 0

    for batch_start in range(0, total, BATCH_SIZE):
        batch = entries_with_coords[batch_start:batch_start + BATCH_SIZE]

        # Build CSV
        csv_buf = io.StringIO()
        writer = csv.writer(csv_buf)
        writer.writerow(["lat", "lon", "osm_id"])
        for entry in batch:
            writer.writerow([entry["lat"], entry["lon"], entry["osm_id"]])

        csv_data = csv_buf.getvalue()

        # Call API (no extra data params - just the CSV file)
        resp = None
        for attempt in range(3):
            try:
                resp = requests.post(
                    GEOCODE_BATCH_URL,
                    files={"data": ("batch.csv", csv_data, "text/csv")},
                    timeout=120,
                )
                if resp.status_code == 200:
                    break
                print(f"    [WARN] Status {resp.status_code}, attempt {attempt+1}")
                time.sleep(2)
            except Exception as e:
                print(f"    [ERROR] Batch geocode attempt {attempt+1}: {e}")
                time.sleep(3)
        if resp is None or resp.status_code != 200:
            print(f"    [FAILED] Batch starting at {batch_start}")
            continue

        # Parse CSV response
        reader = csv.DictReader(io.StringIO(resp.text))
        for row in reader:
            osm_id = row.get("osm_id", "")
            postcode = row.get("result_postcode", "")
            city = row.get("result_city", "")
            status = row.get("result_status", "")

            if osm_id and postcode and status == "ok":
                dept = get_dept_from_postcode(postcode)
                results[osm_id] = {
                    "postcode": postcode,
                    "city": city,
                    "dept": dept,
                }

        processed += len(batch)
        print(f"    Géocodé: {processed:,}/{total:,} ({len(results):,} résolus)")

        # Small delay between batches
        if batch_start + BATCH_SIZE < total:
            time.sleep(1)

    return results


# =============================================================================
# Main processing
# =============================================================================

def load_and_process():
    """Load OSM data, resolve departments, filter to target zones."""
    print(f"1. Chargement de {OSM_FILE}...")
    with open(OSM_FILE, "r", encoding="utf-8") as f:
        elements = json.load(f)
    print(f"  {len(elements):,} entrées chargées")

    # Separate entries by what location data they have
    with_postcode = []
    need_geocoding = []
    no_location = []

    for elem in elements:
        tags = elem.get("tags", {})
        phone = tags.get("phone", "") or tags.get("contact:phone", "")
        name = tags.get("name", "")
        if not phone or not name:
            continue

        pc = tags.get("addr:postcode", "")
        if pc:
            with_postcode.append(elem)
        elif "lat" in elem and "lon" in elem:
            need_geocoding.append(elem)
        else:
            no_location.append(elem)

    print(f"  Avec code postal: {len(with_postcode):,}")
    print(f"  À géocoder (lat/lon): {len(need_geocoding):,}")
    print(f"  Sans localisation: {len(no_location):,}")

    # Phase 1: Filter entries with postcodes
    print(f"\n2. Filtrage par code postal...")
    garages = []
    for elem in with_postcode:
        tags = elem.get("tags", {})
        pc = tags.get("addr:postcode", "")
        dept = get_dept_from_postcode(pc)
        zone = get_zone(dept)
        if zone:
            garages.append(build_garage_entry(elem, pc, dept, zone, tags))

    print(f"  {len(garages):,} garages dans les zones cibles")

    # Phase 2: Batch reverse geocode entries without postcodes
    print(f"\n3. Géocodage inverse des {len(need_geocoding):,} entrées sans code postal...")
    geocode_input = [
        {"lat": e["lat"], "lon": e["lon"], "osm_id": str(e.get("id", ""))}
        for e in need_geocoding
    ]
    geocoded = batch_reverse_geocode(geocode_input)
    print(f"  {len(geocoded):,} entrées géocodées")

    # Filter geocoded entries to target zones
    geocoded_in_zone = 0
    for elem in need_geocoding:
        osm_id = str(elem.get("id", ""))
        if osm_id in geocoded:
            geo = geocoded[osm_id]
            dept = geo["dept"]
            zone = get_zone(dept)
            if zone:
                tags = elem.get("tags", {})
                pc = geo["postcode"]
                city_override = geo.get("city", "")
                entry = build_garage_entry(elem, pc, dept, zone, tags, city_override)
                garages.append(entry)
                geocoded_in_zone += 1

    print(f"  {geocoded_in_zone:,} garages géocodés dans les zones cibles")
    print(f"  Total: {len(garages):,} garages")

    # Deduplicate
    seen = set()
    unique = []
    for g in garages:
        key = g.get("phone_international") or g.get("phone_original")
        if key and key in seen:
            continue
        if key:
            seen.add(key)
        unique.append(g)

    print(f"\n4. Déduplication: {len(garages):,} → {len(unique):,}")
    return unique


def build_garage_entry(elem, postcode, dept, zone, tags, city_override=None):
    """Build a garage dict from an OSM element."""
    phone = tags.get("phone", "") or tags.get("contact:phone", "")
    phone_intl = format_phone_international(phone)

    city = city_override or tags.get("addr:city", "") or tags.get("addr:commune", "")
    street = tags.get("addr:street", "")
    housenumber = tags.get("addr:housenumber", "")
    address = f"{housenumber} {street}".strip() if street else ""

    dept_info = DEPARTMENTS.get(dept, ("", ""))

    return {
        "name": tags.get("name", ""),
        "address": address,
        "postal_code": postcode,
        "city": city,
        "department_num": dept,
        "department_name": dept_info[0],
        "region": dept_info[1],
        "zone": zone,
        "phone_international": phone_intl,
        "phone_original": phone,
        "email": tags.get("email", "") or tags.get("contact:email", ""),
        "website": tags.get("website", "") or tags.get("contact:website", ""),
        "brand": tags.get("brand", ""),
        "opening_hours": tags.get("opening_hours", ""),
        "osm_id": elem.get("id", ""),
        "source": "OpenStreetMap",
    }


# =============================================================================
# Excel output
# =============================================================================

def write_excel(garages, filename):
    """Write garages to a formatted Excel file."""
    print(f"\n5. Écriture de {len(garages):,} garages dans {filename}...")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Garages OSM - Zones Cibles"

    headers = [
        ("N°", 7),
        ("Nom", 45),
        ("Adresse", 40),
        ("Code Postal", 12),
        ("Ville", 25),
        ("Département N°", 14),
        ("Département", 24),
        ("Région", 28),
        ("Zone", 15),
        ("Téléphone (International)", 28),
        ("Téléphone (Original)", 22),
        ("Email", 35),
        ("Site Web", 40),
        ("Marque/Enseigne", 20),
        ("Horaires", 45),
        ("Source", 18),
    ]

    # Styles
    header_font = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill(start_color="1A5276", end_color="1A5276", fill_type="solid")
    header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    thin_border = Border(
        left=Side(style="thin"), right=Side(style="thin"),
        top=Side(style="thin"), bottom=Side(style="thin"),
    )
    data_font = Font(name="Calibri", size=10)
    data_alignment = Alignment(vertical="center", wrap_text=False)

    # Write headers
    for col, (header_text, width) in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header_text)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
        cell.border = thin_border
        ws.column_dimensions[openpyxl.utils.get_column_letter(col)].width = width

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{openpyxl.utils.get_column_letter(len(headers))}1"

    # Sort: Zone → Department → City → Name
    garages.sort(key=lambda x: (
        0 if x.get("zone") == "Top Zone" else 1,
        x.get("department_num", "").zfill(3),
        x.get("city") or "",
        x.get("name") or "",
    ))

    # Zone fills
    top_fill = PatternFill(start_color="E8F8F5", end_color="E8F8F5", fill_type="solid")
    mid_fill = PatternFill(start_color="EBF5FB", end_color="EBF5FB", fill_type="solid")

    for idx, garage in enumerate(garages, 1):
        row = idx + 1
        values = [
            idx,
            garage.get("name", ""),
            garage.get("address", ""),
            garage.get("postal_code", ""),
            garage.get("city", ""),
            garage.get("department_num", ""),
            garage.get("department_name", ""),
            garage.get("region", ""),
            garage.get("zone", ""),
            garage.get("phone_international", ""),
            garage.get("phone_original", ""),
            garage.get("email", ""),
            garage.get("website", ""),
            garage.get("brand", ""),
            garage.get("opening_hours", ""),
            garage.get("source", ""),
        ]

        zone = garage.get("zone", "")
        if zone == "Top Zone":
            row_fill = top_fill if idx % 2 == 0 else PatternFill(fill_type=None)
        else:
            row_fill = mid_fill if idx % 2 == 0 else PatternFill(fill_type=None)

        for col, val in enumerate(values, 1):
            cell = ws.cell(row=row, column=col, value=val)
            cell.font = data_font
            cell.alignment = data_alignment
            cell.border = thin_border
            cell.fill = row_fill

    # Summary sheet
    ws_s = wb.create_sheet("Résumé")
    ws_s.column_dimensions["A"].width = 45
    ws_s.column_dimensions["B"].width = 15

    summary = [
        ("=== STATISTIQUES ===", ""),
        ("Total garages avec téléphone", len(garages)),
        ("Avec email", sum(1 for g in garages if g.get("email"))),
        ("Avec site web", sum(1 for g in garages if g.get("website"))),
        ("Avec marque/enseigne", sum(1 for g in garages if g.get("brand"))),
        ("Avec horaires", sum(1 for g in garages if g.get("opening_hours"))),
        ("", ""),
        ("=== PAR ZONE ===", ""),
        ("  Top Zone", sum(1 for g in garages if g.get("zone") == "Top Zone")),
        ("  Middle Zone", sum(1 for g in garages if g.get("zone") == "Middle Zone")),
        ("", ""),
        ("=== PAR DÉPARTEMENT ===", ""),
    ]

    dept_counts = {}
    for g in garages:
        d = f"{g.get('department_num', '??')} - {g.get('department_name', '??')}"
        dept_counts[d] = dept_counts.get(d, 0) + 1
    for d, count in sorted(dept_counts.items(), key=lambda x: -x[1]):
        summary.append((f"  {d}", count))

    summary_header = Font(name="Calibri", bold=True, size=12, color="1A5276")
    summary_data = Font(name="Calibri", size=11)

    for row_idx, (label, value) in enumerate(summary, 1):
        cell_a = ws_s.cell(row=row_idx, column=1, value=label)
        cell_a.font = summary_header if "===" in str(label) else summary_data
        if value != "":
            cell_b = ws_s.cell(row=row_idx, column=2, value=value)
            cell_b.font = summary_data
            cell_b.alignment = Alignment(horizontal="right")

    wb.save(filename)
    print(f"  Fichier sauvegardé: {filename}")


# =============================================================================
# Main
# =============================================================================

def main():
    print("=" * 70)
    print("  GARAGES OSM AVEC TÉLÉPHONE - Top Zone & Middle Zone")
    print("=" * 70)
    print()

    garages = load_and_process()

    write_excel(garages, OUTPUT_FILE)

    # Final stats
    print("\n" + "=" * 70)
    print("  RÉSUMÉ FINAL")
    print("=" * 70)
    top = sum(1 for g in garages if g.get("zone") == "Top Zone")
    mid = sum(1 for g in garages if g.get("zone") == "Middle Zone")
    emails = sum(1 for g in garages if g.get("email"))
    websites = sum(1 for g in garages if g.get("website"))
    brands = sum(1 for g in garages if g.get("brand"))
    print(f"  Total garages: {len(garages):,}")
    print(f"    Top Zone:    {top:,}")
    print(f"    Middle Zone: {mid:,}")
    print(f"  Avec email:    {emails:,}")
    print(f"  Avec site web: {websites:,}")
    print(f"  Avec marque:   {brands:,}")
    print(f"\n  Fichier: {OUTPUT_FILE}")
    print("\nTerminé!")


if __name__ == "__main__":
    main()
