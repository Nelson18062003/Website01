#!/usr/bin/env python3
"""
Enrich OSM garages with SIRET/SIREN numbers.

Strategy:
1. Load OSM garages from Excel (garages_osm_top_middle_zones.xlsx)
2. For each garage, query recherche-entreprises.api.gouv.fr by name + postal code
3. Validate match with name similarity check
4. Update OSM Excel file with SIRET/SIREN columns
"""

import re
import os
import json
import time
import requests
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

# =============================================================================
# Configuration
# =============================================================================

OSM_EXCEL = "garages_osm_top_middle_zones.xlsx"
OUTPUT_FILE = "garages_osm_top_middle_zones.xlsx"

ENTERPRISE_API_URL = "https://recherche-entreprises.api.gouv.fr/search"
CACHE_FILE = "osm_siret_cache.json"

# =============================================================================
# Name normalization
# =============================================================================

STOP_WORDS = {
    "garage", "garages", "auto", "automobile", "automobiles", "sarl", "sas",
    "eurl", "sa", "srl", "et", "de", "du", "des", "le", "la", "les", "l",
    "un", "une", "en", "au", "aux", "par", "pour", "sur", "chez", "fils",
    "ets", "etablissements", "monsieur", "madame", "mr", "mme", "m",
    "reparation", "entretien", "mecanique", "carrosserie", "service",
    "services", "centre", "atelier", "station", "point", "agent",
    "concessionnaire", "concession", "vente", "france", "international",
    "group", "groupe", "societe", "ste", "société",
}


def normalize_name(name):
    """Normalize a business name for matching."""
    if not name:
        return ""
    n = str(name).lower().strip()
    n = re.sub(r'\([^)]*\)', '', n)
    n = re.sub(r'[^a-z0-9àâäéèêëïîôùûüÿçœæ]', ' ', n)
    n = n.replace('é', 'e').replace('è', 'e').replace('ê', 'e').replace('ë', 'e')
    n = n.replace('à', 'a').replace('â', 'a').replace('ä', 'a')
    n = n.replace('ù', 'u').replace('û', 'u').replace('ü', 'u')
    n = n.replace('î', 'i').replace('ï', 'i')
    n = n.replace('ô', 'o').replace('ç', 'c')
    n = n.replace('œ', 'oe').replace('æ', 'ae').replace('ÿ', 'y')
    n = re.sub(r'\s+', ' ', n).strip()
    return n


def name_tokens(name):
    """Get significant tokens from a normalized name."""
    normalized = normalize_name(name)
    tokens = set(normalized.split())
    tokens = {t for t in tokens if t not in STOP_WORDS and len(t) > 1}
    return tokens


def name_similarity(name1, name2):
    """Token-based similarity (Jaccard index)."""
    tokens1 = name_tokens(name1)
    tokens2 = name_tokens(name2)
    if not tokens1 or not tokens2:
        return 0.0
    intersection = tokens1 & tokens2
    union = tokens1 | tokens2
    return len(intersection) / len(union) if union else 0.0


# =============================================================================
# Cache
# =============================================================================

def load_cache():
    if os.path.exists(CACHE_FILE):
        try:
            with open(CACHE_FILE, "r", encoding="utf-8") as f:
                return json.load(f)
        except (json.JSONDecodeError, IOError):
            pass
    return {}


def save_cache(cache):
    with open(CACHE_FILE, "w", encoding="utf-8") as f:
        json.dump(cache, f, ensure_ascii=False)


# =============================================================================
# Load OSM garages
# =============================================================================

def load_osm_garages(filepath):
    """Load OSM garages from Excel."""
    wb = openpyxl.load_workbook(filepath)
    ws = wb.active

    # Read headers to find column indices
    headers = {}
    for col in range(1, ws.max_column + 1):
        val = ws.cell(row=1, column=col).value
        if val:
            headers[str(val).strip()] = col

    garages = []
    for row in range(2, ws.max_row + 1):
        name = str(ws.cell(row=row, column=headers.get("Nom", 2)).value or "")
        if not name:
            continue
        garages.append({
            "row": row,
            "name": name,
            "address": str(ws.cell(row=row, column=headers.get("Adresse", 3)).value or ""),
            "postal_code": str(ws.cell(row=row, column=headers.get("Code Postal", 4)).value or ""),
            "city": str(ws.cell(row=row, column=headers.get("Ville", 5)).value or ""),
            "dept_num": str(ws.cell(row=row, column=headers.get("Département N°", 6)).value or ""),
        })

    return garages, wb, ws, headers


# =============================================================================
# API lookup
# =============================================================================

def search_siret_api(name, postal_code, city, cache):
    """Search SIRET/SIREN via the French government API."""
    cache_key = f"{name}|{postal_code}|{city}"
    if cache_key in cache:
        cached = cache[cache_key]
        return cached.get("siren", ""), cached.get("siret", ""), cached.get("matched_name", ""), True

    query = str(name).strip()

    # Strategy 1: Search with NAF code filter + postal code
    params = {
        "q": query,
        "per_page": 5,
        "activite_principale": "45.20A,45.20B",
    }
    if postal_code and postal_code not in ("", "None"):
        params["code_postal"] = postal_code

    try:
        resp = requests.get(ENTERPRISE_API_URL, params=params, timeout=15)
        if resp.status_code == 429:
            time.sleep(2)
            resp = requests.get(ENTERPRISE_API_URL, params=params, timeout=15)

        if resp.status_code == 200:
            data = resp.json()
            results = data.get("results", [])
            if results:
                # Check all results for best match
                best_siren, best_siret, best_name, best_score = "", "", "", 0
                for r in results:
                    r_name = r.get("nom_complet", "")
                    score = name_similarity(name, r_name)
                    if score > best_score:
                        best_score = score
                        best_siren = r.get("siren", "")
                        siege = r.get("siege", {})
                        best_siret = siege.get("siret", "")
                        best_name = r_name

                if best_score >= 0.20:
                    cache[cache_key] = {"siren": best_siren, "siret": best_siret, "matched_name": best_name}
                    return best_siren, best_siret, best_name, False

        # Strategy 2: Search without NAF filter (broader search)
        params2 = {
            "q": query,
            "per_page": 5,
        }
        if postal_code and postal_code not in ("", "None"):
            params2["code_postal"] = postal_code

        resp2 = requests.get(ENTERPRISE_API_URL, params=params2, timeout=15)
        if resp2.status_code == 429:
            time.sleep(2)
            resp2 = requests.get(ENTERPRISE_API_URL, params=params2, timeout=15)

        if resp2.status_code == 200:
            data2 = resp2.json()
            results2 = data2.get("results", [])
            if results2:
                best_siren, best_siret, best_name, best_score = "", "", "", 0
                for r in results2:
                    r_name = r.get("nom_complet", "")
                    score = name_similarity(name, r_name)
                    if score > best_score:
                        best_score = score
                        best_siren = r.get("siren", "")
                        siege = r.get("siege", {})
                        best_siret = siege.get("siret", "")
                        best_name = r_name

                if best_score >= 0.25:
                    cache[cache_key] = {"siren": best_siren, "siret": best_siret, "matched_name": best_name}
                    return best_siren, best_siret, best_name, False

    except requests.exceptions.RequestException as e:
        print(f"  [ERROR] API: {e}")

    cache[cache_key] = {"siren": "", "siret": "", "matched_name": ""}
    return "", "", "", False


# =============================================================================
# Excel update
# =============================================================================

def update_excel(garages, wb, ws, headers, output_file):
    """Update the OSM Excel file with SIRET/SIREN columns."""
    max_col = ws.max_column

    # Check if SIREN/SIRET columns already exist
    siren_col = headers.get("SIREN", 0)
    siret_col = headers.get("SIRET", 0)

    if not siren_col:
        siren_col = max_col + 1
        siret_col = max_col + 2

        # Header style
        header_font = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
        header_fill = PatternFill(start_color="1A5276", end_color="1A5276", fill_type="solid")
        header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        thin_border = Border(
            left=Side(style="thin"), right=Side(style="thin"),
            top=Side(style="thin"), bottom=Side(style="thin"),
        )

        cell = ws.cell(row=1, column=siren_col, value="SIREN")
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
        cell.border = thin_border
        ws.column_dimensions[openpyxl.utils.get_column_letter(siren_col)].width = 15

        cell = ws.cell(row=1, column=siret_col, value="SIRET")
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
        cell.border = thin_border
        ws.column_dimensions[openpyxl.utils.get_column_letter(siret_col)].width = 18

    # Write data
    data_font = Font(name="Calibri", size=10)
    data_alignment = Alignment(vertical="center", wrap_text=False)
    thin_border = Border(
        left=Side(style="thin"), right=Side(style="thin"),
        top=Side(style="thin"), bottom=Side(style="thin"),
    )

    updated = 0
    for garage in garages:
        siren = garage.get("siren", "")
        siret = garage.get("siret", "")
        if siren or siret:
            row = garage["row"]
            cell_siren = ws.cell(row=row, column=siren_col, value=siren)
            cell_siren.font = data_font
            cell_siren.alignment = data_alignment
            cell_siren.border = thin_border

            cell_siret = ws.cell(row=row, column=siret_col, value=siret)
            cell_siret.font = data_font
            cell_siret.alignment = data_alignment
            cell_siret.border = thin_border

            updated += 1

    # Update auto filter
    last_col_letter = openpyxl.utils.get_column_letter(siret_col)
    ws.auto_filter.ref = f"A1:{last_col_letter}1"

    # Update summary sheet
    if "Résumé" in wb.sheetnames:
        ws_s = wb["Résumé"]
        next_row = ws_s.max_row + 2
        summary_header = Font(name="Calibri", bold=True, size=12, color="1A5276")
        summary_data = Font(name="Calibri", size=11)

        siren_count = sum(1 for g in garages if g.get("siren"))
        siret_count = sum(1 for g in garages if g.get("siret"))

        ws_s.cell(row=next_row, column=1, value="=== ENRICHISSEMENT SIRET/SIREN ===").font = summary_header
        next_row += 1
        ws_s.cell(row=next_row, column=1, value="Avec SIREN").font = summary_data
        ws_s.cell(row=next_row, column=2, value=siren_count).font = summary_data
        ws_s.cell(row=next_row, column=2).alignment = Alignment(horizontal="right")
        next_row += 1
        ws_s.cell(row=next_row, column=1, value="Avec SIRET").font = summary_data
        ws_s.cell(row=next_row, column=2, value=siret_count).font = summary_data
        ws_s.cell(row=next_row, column=2).alignment = Alignment(horizontal="right")

    wb.save(output_file)
    print(f"\n  {updated:,} garages enrichis avec SIRET/SIREN dans {output_file}")


# =============================================================================
# Main
# =============================================================================

def main():
    print("=" * 70)
    print("  ENRICHISSEMENT SIRET/SIREN DES GARAGES OSM")
    print("=" * 70)
    print()

    # Load OSM garages
    print(f"1. Chargement des garages OSM ({OSM_EXCEL})...")
    osm_garages, wb, ws, headers = load_osm_garages(OSM_EXCEL)
    total = len(osm_garages)
    print(f"  {total:,} garages chargés")

    # Load cache
    cache = load_cache()
    print(f"\n2. Cache: {len(cache)} entrées")

    # Search SIRET/SIREN for each garage
    print(f"\n3. Recherche SIRET/SIREN via l'API entreprises...")
    found_count = 0
    cached_count = 0
    save_interval = 50

    for i, garage in enumerate(osm_garages):
        name = garage["name"]
        postal_code = garage["postal_code"]
        city = garage["city"]

        siren, siret, matched_name, was_cached = search_siret_api(name, postal_code, city, cache)

        if siren or siret:
            garage["siren"] = siren
            garage["siret"] = siret
            found_count += 1

        if was_cached:
            cached_count += 1
        else:
            time.sleep(0.15)  # Rate limit for new API calls

        if (i + 1) % 100 == 0 or (i + 1) == total:
            pct = found_count / (i + 1) * 100
            print(f"  [{i+1:,}/{total:,}] "
                  f"Trouvés: {found_count:,} ({pct:.1f}%) "
                  f"- {cached_count:,} en cache")

        if (i + 1) % save_interval == 0:
            save_cache(cache)

    save_cache(cache)

    # Update Excel
    print(f"\n4. Mise à jour du fichier Excel...")
    update_excel(osm_garages, wb, ws, headers, OUTPUT_FILE)

    # Summary
    print("\n" + "=" * 70)
    print("  RÉSUMÉ FINAL")
    print("=" * 70)
    siren_count = sum(1 for g in osm_garages if g.get("siren"))
    siret_count = sum(1 for g in osm_garages if g.get("siret"))
    print(f"  Total garages OSM: {total:,}")
    print(f"  Avec SIREN: {siren_count:,}")
    print(f"  Avec SIRET: {siret_count:,}")
    print(f"  Sans SIRET/SIREN: {total - siret_count:,}")
    print(f"\n  Fichier: {OUTPUT_FILE}")

    # Cleanup cache
    if os.path.exists(CACHE_FILE):
        os.remove(CACHE_FILE)
        print(f"  Cache supprimé: {CACHE_FILE}")

    print("\nTerminé!")


if __name__ == "__main__":
    main()
