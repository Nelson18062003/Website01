import json
import os

os.chdir('/home/user/Website01')

all_entreprises = []

# --- LILLE ---
for f in ['lille_garages_4520A.json', 'lille_concessions_4511Z.json', 'lille_carrossiers_motos.json']:
    with open(f) as fh:
        data = json.load(fh)
        for e in data:
            e['source_ville'] = 'Lille'
            all_entreprises.append(e)

# --- LYON ---
with open('lyon_auto_moto_entreprises.json') as fh:
    data = json.load(fh)
    for e in data['entreprises']:
        e['source_ville'] = 'Lyon'
        all_entreprises.append(e)

# --- BORDEAUX ---
with open('entreprises_bordeaux.json') as fh:
    data = json.load(fh)
    for e in data['entreprises']:
        e['source_ville'] = 'Bordeaux'
        all_entreprises.append(e)

print(f"Total avant dédoublonnage: {len(all_entreprises)}")

# Normaliser les clés
def normalize(e):
    """Normalize entry to standard keys"""
    nom = e.get('nom') or e.get('nom_complet') or e.get('name', '')
    siret = str(e.get('siret', '')).strip()
    siren = str(e.get('siren', '')).strip()
    adresse = e.get('adresse') or e.get('adresse_complete') or ''
    code_postal = str(e.get('code_postal', '')).strip()
    ville = e.get('ville') or e.get('libelle_commune') or ''
    code_naf = e.get('code_naf') or e.get('activite_principale') or ''
    libelle = e.get('libelle_activite') or e.get('libelle_naf') or ''
    date_creation = e.get('date_creation') or e.get('date_creation_entreprise') or ''
    statut_raw = e.get('statut') or e.get('etat_administratif') or ''
    if statut_raw in ('A', 'Actif', 'actif'):
        statut = 'Actif'
    elif statut_raw in ('C', 'Cessé', 'cessé', 'Fermé', 'fermé', 'Cesse'):
        statut = 'Cessé'
    else:
        statut = statut_raw
    effectif = e.get('effectif') or e.get('tranche_effectif') or ''
    categorie = e.get('categorie') or ''
    source_ville = e.get('source_ville', '')
    
    return {
        'nom': nom,
        'siret': siret,
        'siren': siren,
        'adresse': adresse,
        'code_postal': code_postal,
        'ville': ville,
        'code_naf': code_naf,
        'libelle_activite': libelle,
        'date_creation': date_creation,
        'statut': statut,
        'effectif': str(effectif),
        'categorie': categorie,
        'zone_recherche': source_ville
    }

normalized = [normalize(e) for e in all_entreprises]

# Dédoublonnage par SIRET
seen = set()
unique = []
for e in normalized:
    key = e['siret']
    if key and key not in seen:
        seen.add(key)
        unique.append(e)
    elif not key:
        unique.append(e)

print(f"Total après dédoublonnage par SIRET: {len(unique)}")

# Filtrer uniquement les actifs
actifs = [e for e in unique if e['statut'] == 'Actif']
print(f"Total entreprises actives: {len(actifs)}")

# Stats par ville et catégorie
from collections import Counter
print("\n--- Par zone de recherche ---")
for ville, count in Counter(e['zone_recherche'] for e in actifs).most_common():
    print(f"  {ville}: {count}")

print("\n--- Par catégorie (actifs) ---")
for cat, count in Counter(e['categorie'] for e in actifs).most_common():
    print(f"  {cat}: {count}")

# Export JSON
with open('garages_lille_lyon_bordeaux.json', 'w', encoding='utf-8') as f:
    json.dump(actifs, f, ensure_ascii=False, indent=2)
print(f"\nJSON exporté: garages_lille_lyon_bordeaux.json")

# Export Excel
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

wb = Workbook()
ws = wb.active
ws.title = "Entreprises Auto-Moto"

# Headers
headers = ['Nom', 'SIRET', 'SIREN', 'Adresse', 'Code Postal', 'Ville', 
           'Code NAF', 'Activité', 'Date Création', 'Statut', 'Effectif', 
           'Catégorie', 'Zone Recherche']

header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
header_font = Font(color="FFFFFF", bold=True, size=11)
thin_border = Border(
    left=Side(style='thin'),
    right=Side(style='thin'),
    top=Side(style='thin'),
    bottom=Side(style='thin')
)

for col, h in enumerate(headers, 1):
    cell = ws.cell(row=1, column=col, value=h)
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = Alignment(horizontal='center')
    cell.border = thin_border

# Data
keys = ['nom', 'siret', 'siren', 'adresse', 'code_postal', 'ville',
        'code_naf', 'libelle_activite', 'date_creation', 'statut', 'effectif',
        'categorie', 'zone_recherche']

for row_idx, e in enumerate(actifs, 2):
    for col_idx, key in enumerate(keys, 1):
        val = e.get(key, '')
        cell = ws.cell(row=row_idx, column=col_idx, value=val)
        cell.border = thin_border
        # SIRET/SIREN as text
        if key in ('siret', 'siren', 'code_postal'):
            cell.number_format = '@'

# Auto-width
for col in ws.columns:
    max_length = 0
    for cell in col:
        if cell.value:
            max_length = max(max_length, len(str(cell.value)))
    ws.column_dimensions[col[0].column_letter].width = min(max_length + 2, 50)

# Freeze header
ws.freeze_panes = 'A2'
ws.auto_filter.ref = ws.dimensions

wb.save('garages_lille_lyon_bordeaux.xlsx')
print(f"Excel exporté: garages_lille_lyon_bordeaux.xlsx")
