#!/usr/bin/env python3
"""Merge all scraped data, deduplicate, enrich, and export to XLSX."""
import os
import re
import json
import time
import pandas as pd
import requests
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from scraper_utils import (
    DATA_DIR, NAF_CODES, ZONES, setup_logging, load_json,
    normalize_name, format_phone, format_postal_code,
    categorize_by_naf_and_name, save_json
)

logger = setup_logging("merge_export")

OUTPUT_FILE = os.path.join(os.path.dirname(__file__), "structures-medicales.xlsx")

COLUMNS = [
    "raison_sociale", "categorie", "code_naf", "libelle_naf",
    "telephone", "adresse", "code_postal", "ville",
    "siret", "siren", "site_web", "email",
    "effectif", "date_creation", "source"
]

LILLE_POSTCODES = {"59000", "59800", "59100", "59200", "59491", "59130", "59110",
                   "59160", "59170", "59260", "59650", "59700", "59790", "59810"}


def load_all_sources():
    """Load all intermediate JSON files."""
    all_records = []

    # API Entreprises
    data = load_json("api_entreprises.json")
    logger.info(f"Loaded {len(data)} from API Entreprises")
    all_records.extend(data)

    # FINESS
    data = load_json("finess.json")
    logger.info(f"Loaded {len(data)} from FINESS")
    all_records.extend(data)

    # OSM
    data = load_json("osm.json")
    logger.info(f"Loaded {len(data)} from OSM")
    all_records.extend(data)

    # Pages Jaunes (multiple files)
    for fname in os.listdir(DATA_DIR):
        if fname.startswith("pagesjaunes_") and fname.endswith(".json"):
            data = load_json(fname)
            logger.info(f"Loaded {len(data)} from {fname}")
            all_records.extend(data)

    # Annuaire Santé (multiple files)
    for fname in os.listdir(DATA_DIR):
        if fname.startswith("annuaire_sante_") and fname.endswith(".json"):
            data = load_json(fname)
            logger.info(f"Loaded {len(data)} from {fname}")
            all_records.extend(data)

    logger.info(f"Total raw records from all sources: {len(all_records)}")
    return all_records


def deduplicate(records):
    """Deduplicate records. Primary key: SIRET. Secondary: normalized name + postal code."""
    df = pd.DataFrame(records)

    # Ensure all columns exist
    for col in COLUMNS:
        if col not in df.columns:
            df[col] = ""

    # Fill NaN with empty string
    df = df.fillna("")

    # Convert all to string
    for col in df.columns:
        df[col] = df[col].astype(str)

    # Clean up SIRET
    df["siret"] = df["siret"].apply(lambda x: re.sub(r'[^\d]', '', str(x)))
    df["siret"] = df["siret"].apply(lambda x: x if len(x) == 14 else "")

    # Extract SIREN from SIRET if missing
    df["siren"] = df.apply(lambda row: row["siret"][:9] if row["siret"] and len(row["siret"]) == 14 and not row["siren"] else row["siren"], axis=1)

    # Normalize name for dedup
    df["_norm_name"] = df["raison_sociale"].apply(normalize_name)
    df["_norm_cp"] = df["code_postal"].apply(format_postal_code)

    # Score each record by completeness
    def completeness_score(row):
        score = 0
        if row["siret"]: score += 5
        if row["telephone"]: score += 3
        if row["adresse"]: score += 2
        if row["email"]: score += 2
        if row["site_web"]: score += 1
        if row["effectif"]: score += 1
        if row["date_creation"]: score += 1
        if row["code_naf"]: score += 2
        return score

    df["_score"] = df.apply(completeness_score, axis=1)

    # Sort by score descending so best records come first
    df = df.sort_values("_score", ascending=False)

    # Dedup by SIRET first
    seen_sirets = {}
    seen_names = {}
    result_indices = []

    for idx, row in df.iterrows():
        siret = row["siret"]
        norm_key = f"{row['_norm_name']}_{row['_norm_cp']}"

        if siret and siret in seen_sirets:
            # Merge non-empty fields into existing record
            existing_idx = seen_sirets[siret]
            for col in COLUMNS:
                if col == "source":
                    existing_src = df.at[existing_idx, "source"]
                    new_src = row["source"]
                    if new_src and new_src not in existing_src:
                        df.at[existing_idx, "source"] = f"{existing_src}; {new_src}"
                elif not df.at[existing_idx, col] and row[col]:
                    df.at[existing_idx, col] = row[col]
            continue

        if norm_key and norm_key in seen_names and not siret:
            # Merge by name
            existing_idx = seen_names[norm_key]
            for col in COLUMNS:
                if col == "source":
                    existing_src = df.at[existing_idx, "source"]
                    new_src = row["source"]
                    if new_src and new_src not in existing_src:
                        df.at[existing_idx, "source"] = f"{existing_src}; {new_src}"
                elif not df.at[existing_idx, col] and row[col]:
                    df.at[existing_idx, col] = row[col]
            continue

        # New record
        result_indices.append(idx)
        if siret:
            seen_sirets[siret] = idx
        if norm_key:
            seen_names[norm_key] = idx

    df_dedup = df.loc[result_indices].copy()

    # Drop temp columns
    df_dedup = df_dedup.drop(columns=["_norm_name", "_norm_cp", "_score"], errors='ignore')

    logger.info(f"After deduplication: {len(df_dedup)} unique records (from {len(df)} raw)")
    return df_dedup


def enrich_siret(df, max_enrichments=2000):
    """Try to enrich missing SIRETs via API Recherche Entreprises."""
    missing = df[df["siret"] == ""].copy()
    if len(missing) == 0:
        logger.info("No missing SIRETs to enrich")
        return df

    logger.info(f"Attempting SIRET enrichment for {min(len(missing), max_enrichments)} records...")
    enriched = 0

    for idx, row in missing.head(max_enrichments).iterrows():
        name = row["raison_sociale"]
        cp = row["code_postal"]
        if not name or not cp:
            continue

        try:
            params = {
                "q": name,
                "code_postal": cp,
                "etat_administratif": "A",
                "page": 1,
                "per_page": 1,
            }
            r = requests.get("https://recherche-entreprises.api.gouv.fr/search",
                           params=params, timeout=15)
            if r.status_code == 200:
                data = r.json()
                results = data.get("results", [])
                if results:
                    ent = results[0]
                    siege = ent.get("siege", {})
                    siret = siege.get("siret", "")
                    if siret and len(re.sub(r'[^\d]', '', siret)) == 14:
                        df.at[idx, "siret"] = re.sub(r'[^\d]', '', siret)
                        df.at[idx, "siren"] = re.sub(r'[^\d]', '', siret)[:9]
                        if not df.at[idx, "code_naf"]:
                            naf = siege.get("activite_principale", "")
                            df.at[idx, "code_naf"] = naf
                        if not df.at[idx, "effectif"]:
                            df.at[idx, "effectif"] = siege.get("tranche_effectif_salarie", "") or ent.get("tranche_effectif_salarie", "")
                        if not df.at[idx, "date_creation"]:
                            df.at[idx, "date_creation"] = ent.get("date_creation", "")
                        enriched += 1
            time.sleep(0.5)
        except Exception as e:
            continue

    logger.info(f"SIRET enrichment: {enriched} records enriched")
    return df


def cleanup(df):
    """Clean and standardize the dataframe."""
    # Format phones
    df["telephone"] = df["telephone"].apply(format_phone)

    # Format postal codes
    df["code_postal"] = df["code_postal"].apply(format_postal_code)

    # Standardize NAF codes (with dot format)
    def fix_naf(naf):
        naf = str(naf).strip()
        if len(naf) == 5 and naf[2] != '.':
            return naf[:2] + '.' + naf[2:]
        return naf

    df["code_naf"] = df["code_naf"].apply(fix_naf)

    # Fill libelle_naf from code
    for idx, row in df.iterrows():
        if not row["libelle_naf"] and row["code_naf"] in NAF_CODES:
            df.at[idx, "libelle_naf"] = NAF_CODES[row["code_naf"]]

    # Re-categorize if needed
    for idx, row in df.iterrows():
        if not row["categorie"] or row["categorie"] == "Structure médicale":
            if row["code_naf"]:
                df.at[idx, "categorie"] = categorize_by_naf_and_name(row["code_naf"], row["raison_sociale"])

    # Capitalize city names
    df["ville"] = df["ville"].str.upper()

    # Remove empty raison_sociale
    df = df[df["raison_sociale"].str.strip() != ""]

    # Remove 'nan' strings
    for col in df.columns:
        df[col] = df[col].replace('nan', '')

    # Keep only standard columns
    df = df[[c for c in COLUMNS if c in df.columns]]

    return df


def get_zone(cp):
    """Determine zone from postal code."""
    cp = str(cp)
    if cp.startswith("75"):
        return "paris"
    if cp.startswith("13"):
        return "marseille"
    if cp.startswith("59"):
        return "lille"
    return "autre"


def export_xlsx(df):
    """Export to formatted XLSX with multiple sheets."""
    logger.info(f"Exporting {len(df)} records to {OUTPUT_FILE}")

    # Add zone column for filtering
    df["_zone"] = df["code_postal"].apply(get_zone)

    df_paris = df[df["_zone"] == "paris"].drop(columns=["_zone"])
    df_marseille = df[df["_zone"] == "marseille"].drop(columns=["_zone"])
    df_lille = df[df["_zone"] == "lille"].drop(columns=["_zone"])
    df_all = df.drop(columns=["_zone"])

    # Write with pandas first
    with pd.ExcelWriter(OUTPUT_FILE, engine='openpyxl') as writer:
        df_all.to_excel(writer, sheet_name="Toutes les structures", index=False)
        df_paris.to_excel(writer, sheet_name="Paris", index=False)
        df_marseille.to_excel(writer, sheet_name="Marseille", index=False)
        df_lille.to_excel(writer, sheet_name="Lille", index=False)

        # Statistics sheet
        stats = build_statistics(df_all, df_paris, df_marseille, df_lille)
        stats.to_excel(writer, sheet_name="Statistiques", index=False)

    # Format with openpyxl
    format_xlsx(OUTPUT_FILE)
    logger.info(f"Excel file saved: {OUTPUT_FILE}")


def build_statistics(df_all, df_paris, df_marseille, df_lille):
    """Build statistics DataFrame."""
    stats_rows = []

    # Total
    stats_rows.append({"Statistique": "RÉSUMÉ GÉNÉRAL", "Valeur": ""})
    stats_rows.append({"Statistique": "Total structures", "Valeur": len(df_all)})
    stats_rows.append({"Statistique": "", "Valeur": ""})

    # By city
    stats_rows.append({"Statistique": "RÉPARTITION PAR VILLE", "Valeur": ""})
    stats_rows.append({"Statistique": "Paris (75)", "Valeur": len(df_paris)})
    stats_rows.append({"Statistique": "Marseille (13)", "Valeur": len(df_marseille)})
    stats_rows.append({"Statistique": "Lille et agglo (59)", "Valeur": len(df_lille)})
    stats_rows.append({"Statistique": "", "Valeur": ""})

    # By category
    stats_rows.append({"Statistique": "RÉPARTITION PAR CATÉGORIE", "Valeur": ""})
    if "categorie" in df_all.columns:
        cat_counts = df_all["categorie"].value_counts()
        for cat, count in cat_counts.items():
            if cat:
                stats_rows.append({"Statistique": f"  {cat}", "Valeur": count})
    stats_rows.append({"Statistique": "", "Valeur": ""})

    # By NAF code
    stats_rows.append({"Statistique": "RÉPARTITION PAR CODE NAF", "Valeur": ""})
    if "code_naf" in df_all.columns:
        naf_counts = df_all["code_naf"].value_counts()
        for naf, count in naf_counts.items():
            if naf:
                label = NAF_CODES.get(naf, "")
                stats_rows.append({"Statistique": f"  {naf} — {label}", "Valeur": count})
    stats_rows.append({"Statistique": "", "Valeur": ""})

    # Completeness rates
    stats_rows.append({"Statistique": "TAUX DE COMPLÉTUDE", "Valeur": ""})
    total = len(df_all) if len(df_all) > 0 else 1
    siret_pct = round(len(df_all[df_all["siret"] != ""]) / total * 100, 1)
    tel_pct = round(len(df_all[df_all["telephone"] != ""]) / total * 100, 1)
    email_pct = round(len(df_all[df_all["email"] != ""]) / total * 100, 1)
    web_pct = round(len(df_all[df_all["site_web"] != ""]) / total * 100, 1)
    addr_pct = round(len(df_all[df_all["adresse"] != ""]) / total * 100, 1)

    stats_rows.append({"Statistique": "  SIRET", "Valeur": f"{siret_pct}%"})
    stats_rows.append({"Statistique": "  Téléphone", "Valeur": f"{tel_pct}%"})
    stats_rows.append({"Statistique": "  Email", "Valeur": f"{email_pct}%"})
    stats_rows.append({"Statistique": "  Site web", "Valeur": f"{web_pct}%"})
    stats_rows.append({"Statistique": "  Adresse", "Valeur": f"{addr_pct}%"})
    stats_rows.append({"Statistique": "", "Valeur": ""})

    # Sources
    stats_rows.append({"Statistique": "SOURCES UTILISÉES", "Valeur": ""})
    if "source" in df_all.columns:
        # Count primary source for each record
        source_counts = {}
        for src in df_all["source"]:
            primary = str(src).split(";")[0].strip()
            source_counts[primary] = source_counts.get(primary, 0) + 1
        for src, count in sorted(source_counts.items(), key=lambda x: -x[1]):
            if src:
                stats_rows.append({"Statistique": f"  {src}", "Valeur": count})

    return pd.DataFrame(stats_rows)


def format_xlsx(filepath):
    """Format the XLSX file with openpyxl."""
    wb = load_workbook(filepath)

    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill(start_color="2E75B6", end_color="2E75B6", fill_type="solid")
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    thin_border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]

        # Format headers
        for cell in ws[1]:
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_align
            cell.border = thin_border

        # Auto-width columns
        for col_idx, col_cells in enumerate(ws.columns, 1):
            max_len = 0
            for cell in col_cells:
                try:
                    cell_len = len(str(cell.value or ""))
                    max_len = max(max_len, min(cell_len, 50))
                except:
                    pass
            col_letter = get_column_letter(col_idx)
            ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

        # Auto-filter on data sheets
        if sheet_name != "Statistiques" and ws.max_row > 1:
            ws.auto_filter.ref = f"A1:{get_column_letter(ws.max_column)}{ws.max_row}"

        # Freeze top row
        ws.freeze_panes = "A2"

    # Special formatting for Statistics sheet
    if "Statistiques" in wb.sheetnames:
        ws = wb["Statistiques"]
        section_font = Font(bold=True, size=12, color="2E75B6")
        for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
            val = str(row[0].value or "")
            if val.isupper() and val:
                row[0].font = section_font

    wb.save(filepath)
    logger.info("XLSX formatting applied")


def print_summary(df):
    """Print a summary to the terminal."""
    total = len(df)
    print("\n" + "=" * 70)
    print("  RÉSUMÉ — STRUCTURES MÉDICALES ET DE SANTÉ")
    print("=" * 70)
    print(f"\n  Total structures : {total:,}")
    print()

    # By zone
    df_temp = df.copy()
    df_temp["_zone"] = df_temp["code_postal"].apply(get_zone)
    print("  Répartition par ville :")
    for zone, label in [("paris", "Paris (75)"), ("marseille", "Marseille (13)"), ("lille", "Lille (59)")]:
        count = len(df_temp[df_temp["_zone"] == zone])
        print(f"    {label:30s} : {count:>6,}")

    # By category
    print("\n  Répartition par catégorie :")
    cat_counts = df["categorie"].value_counts().head(20)
    for cat, count in cat_counts.items():
        if cat:
            print(f"    {cat:45s} : {count:>6,}")

    # By NAF code
    print("\n  Répartition par code NAF :")
    naf_counts = df["code_naf"].value_counts().head(10)
    for naf, count in naf_counts.items():
        if naf:
            label = NAF_CODES.get(naf, "")
            print(f"    {naf} — {label:45s} : {count:>6,}")

    # Completeness
    t = max(total, 1)
    print("\n  Taux de complétude :")
    print(f"    SIRET      : {len(df[df['siret'] != '']) / t * 100:5.1f}%")
    print(f"    Téléphone  : {len(df[df['telephone'] != '']) / t * 100:5.1f}%")
    print(f"    Email      : {len(df[df['email'] != '']) / t * 100:5.1f}%")
    print(f"    Site web   : {len(df[df['site_web'] != '']) / t * 100:5.1f}%")
    print(f"    Adresse    : {len(df[df['adresse'] != '']) / t * 100:5.1f}%")

    # Sources
    print("\n  Sources utilisées :")
    source_counts = {}
    for src in df["source"]:
        primary = str(src).split(";")[0].strip()
        source_counts[primary] = source_counts.get(primary, 0) + 1
    for src, count in sorted(source_counts.items(), key=lambda x: -x[1]):
        if src:
            print(f"    {src:40s} : {count:>6,}")

    print("\n" + "=" * 70)
    print(f"  Fichier exporté : structures-medicales.xlsx")
    print("=" * 70 + "\n")


def main():
    logger.info("=== Starting merge and export ===")

    # Load all sources
    records = load_all_sources()
    if not records:
        logger.error("No data found! Make sure scrapers have run first.")
        return

    # Deduplicate
    df = deduplicate(records)

    # Enrich missing SIRETs
    df = enrich_siret(df)

    # Cleanup
    df = cleanup(df)

    # Export
    export_xlsx(df)

    # Print summary
    print_summary(df)

    logger.info("=== Merge and export complete ===")


if __name__ == "__main__":
    main()
