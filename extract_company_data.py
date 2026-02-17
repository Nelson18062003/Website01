#!/usr/bin/env python3
"""
Script to enrich company data from Pages Jaunes with:
- SIRET / SIREN numbers (from recherche-entreprises.api.gouv.fr)
- Department number and name (from postal code)
"""

import openpyxl
import requests
import time
import sys
import os
import json

# Mapping of department numbers to names for Ile-de-France
DEPARTMENTS_IDF = {
    "75": "Paris",
    "77": "Seine-et-Marne",
    "78": "Yvelines",
    "91": "Essonne",
    "92": "Hauts-de-Seine",
    "93": "Seine-Saint-Denis",
    "94": "Val-de-Marne",
    "95": "Val-d'Oise",
}

API_URL = "https://recherche-entreprises.api.gouv.fr/search"

# Cache file to resume if interrupted
CACHE_FILE = "siret_cache.json"


def get_department_info(postal_code):
    """Extract department number and name from postal code."""
    if postal_code is None:
        return "", ""
    pc = str(postal_code).strip()
    if len(pc) >= 2:
        dept_num = pc[:2]
        dept_name = DEPARTMENTS_IDF.get(dept_num, "")
        return dept_num, dept_name
    return "", ""


def search_company(company_name, postal_code, city):
    """Search for company SIRET/SIREN via the French government API."""
    if not company_name:
        return "", ""

    # Build query
    query = str(company_name).strip()
    params = {
        "q": query,
        "per_page": 5,
    }

    # Add postal code filter if available
    if postal_code:
        pc = str(postal_code).strip()
        if pc and pc != "None":
            params["code_postal"] = pc

    try:
        resp = requests.get(API_URL, params=params, timeout=15)
        if resp.status_code == 429:
            # Rate limited, wait and retry
            time.sleep(2)
            resp = requests.get(API_URL, params=params, timeout=15)

        if resp.status_code == 200:
            data = resp.json()
            results = data.get("results", [])
            if results:
                # Take the first (best) match
                best = results[0]
                siren = best.get("siren", "")
                # Get the matching siege (headquarters) for SIRET
                siege = best.get("siege", {})
                siret = siege.get("siret", "")
                matched_name = best.get("nom_complet", "")
                return siren, siret, matched_name

        # If postal code search failed, try with city
        if city and str(city).strip() != "None":
            params_city = {
                "q": query,
                "per_page": 5,
                "commune": str(city).strip(),
            }
            # Remove commune param, use q with city instead
            params_city = {
                "q": f"{query} {city}",
                "per_page": 5,
            }
            resp2 = requests.get(API_URL, params=params_city, timeout=15)
            if resp2.status_code == 200:
                data2 = resp2.json()
                results2 = data2.get("results", [])
                if results2:
                    best = results2[0]
                    siren = best.get("siren", "")
                    siege = best.get("siege", {})
                    siret = siege.get("siret", "")
                    matched_name = best.get("nom_complet", "")
                    return siren, siret, matched_name

    except requests.exceptions.RequestException as e:
        print(f"  [ERROR] API request failed: {e}")

    return "", "", ""


def load_cache():
    """Load cached results to allow resuming."""
    if os.path.exists(CACHE_FILE):
        with open(CACHE_FILE, "r", encoding="utf-8") as f:
            return json.load(f)
    return {}


def save_cache(cache):
    """Save cache to disk."""
    with open(CACHE_FILE, "w", encoding="utf-8") as f:
        json.dump(cache, f, ensure_ascii=False)


def main():
    input_file = "pagesjaunes_Centres médicaux_Ile-de-France.xlsx"
    output_file = "pagesjaunes_Centres_medicaux_IDF_enriched.xlsx"

    print(f"Loading {input_file}...")
    wb = openpyxl.load_workbook(input_file)
    ws = wb.active
    total_rows = ws.max_row - 1  # excluding header
    print(f"Found {total_rows} companies to process.")

    # Load cache
    cache = load_cache()
    print(f"Cache contains {len(cache)} entries.")

    # Create output workbook
    wb_out = openpyxl.Workbook()
    ws_out = wb_out.active
    ws_out.title = "Enriched Data"

    # Write headers
    headers = [
        "company_name",
        "street_address",
        "postal_code",
        "city",
        "region",
        "country",
        "phone_number",
        "phone_format_fr",
        "departement_numero",
        "departement_nom",
        "SIREN",
        "SIRET",
        "nom_api",
    ]
    for col, header in enumerate(headers, 1):
        ws_out.cell(row=1, column=col, value=header)

    found_count = 0
    not_found_count = 0

    for row_idx in range(2, ws.max_row + 1):
        company_name = ws.cell(row=row_idx, column=1).value
        street_address = ws.cell(row=row_idx, column=2).value
        postal_code = ws.cell(row=row_idx, column=3).value
        city = ws.cell(row=row_idx, column=4).value
        region = ws.cell(row=row_idx, column=5).value
        country = ws.cell(row=row_idx, column=6).value
        phone_number = ws.cell(row=row_idx, column=7).value
        phone_format_fr = ws.cell(row=row_idx, column=8).value

        # Department info from postal code
        dept_num, dept_name = get_department_info(postal_code)

        # Cache key
        cache_key = f"{company_name}|{postal_code}|{city}"

        if cache_key in cache:
            siren, siret, matched_name = cache[cache_key]
        else:
            # Query API for SIRET/SIREN
            siren, siret, matched_name = search_company(company_name, postal_code, city)
            cache[cache_key] = [siren, siret, matched_name]

            # Small delay to respect rate limits
            time.sleep(0.15)

        if siren:
            found_count += 1
        else:
            not_found_count += 1

        current = row_idx - 1
        if current % 50 == 0 or current == total_rows:
            print(f"  Progress: {current}/{total_rows} | Found: {found_count} | Not found: {not_found_count}")
            save_cache(cache)

        # Write row to output
        out_row = row_idx
        ws_out.cell(row=out_row, column=1, value=company_name)
        ws_out.cell(row=out_row, column=2, value=street_address)
        ws_out.cell(row=out_row, column=3, value=postal_code)
        ws_out.cell(row=out_row, column=4, value=city)
        ws_out.cell(row=out_row, column=5, value=region)
        ws_out.cell(row=out_row, column=6, value=country)
        ws_out.cell(row=out_row, column=7, value=phone_number)
        ws_out.cell(row=out_row, column=8, value=phone_format_fr)
        ws_out.cell(row=out_row, column=9, value=dept_num)
        ws_out.cell(row=out_row, column=10, value=dept_name)
        ws_out.cell(row=out_row, column=11, value=siren)
        ws_out.cell(row=out_row, column=12, value=siret)
        ws_out.cell(row=out_row, column=13, value=matched_name)

    # Save final cache
    save_cache(cache)

    # Save output file
    print(f"\nSaving {output_file}...")
    wb_out.save(output_file)
    print(f"Done! Output saved to {output_file}")
    print(f"Total: {total_rows} | SIREN found: {found_count} | Not found: {not_found_count}")

    # Cleanup cache file
    if os.path.exists(CACHE_FILE):
        os.remove(CACHE_FILE)


if __name__ == "__main__":
    main()
