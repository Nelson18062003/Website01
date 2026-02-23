#!/usr/bin/env python3
"""
Fast export of PJ garages to Excel with SIRET enrichment using multithreading.
"""

import json
import os
import re
import time
from concurrent.futures import ThreadPoolExecutor, as_completed
from threading import Lock

import openpyxl
import requests

CACHE_DIR = "pagesjaunes_garages_cache"
SIRET_CACHE_FILE = "pj_garages_siret_cache.json"
OUTPUT_FILE = "pagesjaunes_garages_france.xlsx"
API_URL = "https://recherche-entreprises.api.gouv.fr/search"
WORKERS = 10  # parallel threads

ALL_DEPTS = [
    "01", "04", "13", "14", "16", "21", "22", "23", "25", "26",
    "27", "28", "29", "30", "31", "32", "33", "34", "35", "37",
    "38", "39", "42", "44", "45", "46", "49", "50", "53", "56",
    "57", "59", "60", "61", "62", "64", "65", "66", "67", "68",
    "69", "70", "71", "73", "74", "75", "76", "77", "80", "83",
    "85", "87", "88", "91", "92", "93", "94", "95", "2B",
]

DEPT_NAMES = {
    "01": "Ain", "04": "Alpes-de-Haute-Provence",
    "13": "Bouches-du-Rhône", "14": "Calvados", "16": "Charente",
    "21": "Côte-d'Or", "22": "Côtes-d'Armor", "23": "Creuse",
    "25": "Doubs", "26": "Drôme", "27": "Eure", "28": "Eure-et-Loir",
    "29": "Finistère", "30": "Gard", "31": "Haute-Garonne", "32": "Gers",
    "33": "Gironde", "34": "Hérault", "35": "Ille-et-Vilaine",
    "37": "Indre-et-Loire", "38": "Isère", "39": "Jura",
    "42": "Loire", "44": "Loire-Atlantique", "45": "Loiret", "46": "Lot",
    "49": "Maine-et-Loire", "50": "Manche", "53": "Mayenne",
    "56": "Morbihan", "57": "Moselle", "59": "Nord", "60": "Oise",
    "61": "Orne", "62": "Pas-de-Calais", "64": "Pyrénées-Atlantiques",
    "65": "Hautes-Pyrénées", "66": "Pyrénées-Orientales", "67": "Bas-Rhin",
    "68": "Haut-Rhin", "69": "Rhône", "70": "Haute-Saône",
    "71": "Saône-et-Loire", "73": "Savoie", "74": "Haute-Savoie",
    "75": "Paris", "76": "Seine-Maritime", "77": "Seine-et-Marne",
    "80": "Somme", "83": "Var", "85": "Vendée", "87": "Haute-Vienne",
    "88": "Vosges", "91": "Essonne", "92": "Hauts-de-Seine",
    "93": "Seine-Saint-Denis", "94": "Val-de-Marne", "95": "Val-d'Oise",
    "2B": "Haute-Corse",
}

REGION_MAP = {
    "01": "Auvergne-Rhône-Alpes", "26": "Auvergne-Rhône-Alpes",
    "38": "Auvergne-Rhône-Alpes", "42": "Auvergne-Rhône-Alpes",
    "69": "Auvergne-Rhône-Alpes", "73": "Auvergne-Rhône-Alpes",
    "74": "Auvergne-Rhône-Alpes",
    "21": "Bourgogne-Franche-Comté", "25": "Bourgogne-Franche-Comté",
    "39": "Bourgogne-Franche-Comté", "70": "Bourgogne-Franche-Comté",
    "71": "Bourgogne-Franche-Comté",
    "22": "Bretagne", "29": "Bretagne", "35": "Bretagne", "56": "Bretagne",
    "28": "Centre-Val de Loire", "37": "Centre-Val de Loire", "45": "Centre-Val de Loire",
    "2B": "Corse",
    "57": "Grand Est", "67": "Grand Est", "68": "Grand Est", "88": "Grand Est",
    "59": "Hauts-de-France", "60": "Hauts-de-France",
    "62": "Hauts-de-France", "80": "Hauts-de-France",
    "75": "Île-de-France", "77": "Île-de-France", "91": "Île-de-France",
    "92": "Île-de-France", "93": "Île-de-France",
    "94": "Île-de-France", "95": "Île-de-France",
    "14": "Normandie", "27": "Normandie", "50": "Normandie",
    "61": "Normandie", "76": "Normandie",
    "16": "Nouvelle-Aquitaine", "23": "Nouvelle-Aquitaine",
    "33": "Nouvelle-Aquitaine", "64": "Nouvelle-Aquitaine",
    "85": "Pays de la Loire", "87": "Nouvelle-Aquitaine",
    "30": "Occitanie", "31": "Occitanie", "32": "Occitanie",
    "34": "Occitanie", "46": "Occitanie", "65": "Occitanie",
    "66": "Occitanie",
    "44": "Pays de la Loire", "49": "Pays de la Loire",
    "53": "Pays de la Loire",
    "04": "Provence-Alpes-Côte d'Azur", "13": "Provence-Alpes-Côte d'Azur",
    "83": "Provence-Alpes-Côte d'Azur",
}

cache = {}
cache_lock = Lock()
counter = {"done": 0, "found": 0, "not_found": 0, "cached": 0}
counter_lock = Lock()


def load_all_garages():
    all_garages = []
    target = set(ALL_DEPTS)
    files = sorted(f for f in os.listdir(CACHE_DIR) if f.endswith('.json'))
    for fname in files:
        dept = fname.replace("dept_", "").replace(".json", "")
        if dept not in target:
            print(f"  {fname}: SKIP (not in target departments)")
            continue
        with open(os.path.join(CACHE_DIR, fname), "r", encoding="utf-8") as f:
            garages = json.load(f)
        for g in garages:
            g["dept_num"] = dept
            g["dept_name"] = DEPT_NAMES.get(dept, "")
            g["region"] = REGION_MAP.get(dept, "")
        all_garages.extend(garages)
        print(f"  {fname}: {len(garages)} garages")
    return all_garages


def search_siret(name, postal_code, city):
    if not name:
        return "", "", "", ""
    query = re.sub(r'\s*\(.*?\)\s*', ' ', str(name).strip()).strip()
    params = {"q": query, "per_page": 5}
    if postal_code and str(postal_code).strip() not in ("", "None"):
        params["code_postal"] = str(postal_code).strip()

    session = requests.Session()
    for attempt in range(3):
        try:
            resp = session.get(API_URL, params=params, timeout=15)
            if resp.status_code == 429:
                time.sleep(2 * (attempt + 1))
                continue
            if resp.status_code == 200:
                results = resp.json().get("results", [])
                if results:
                    best = results[0]
                    siege = best.get("siege", {})
                    return (best.get("siren", ""), siege.get("siret", ""),
                            best.get("nom_complet", ""), best.get("activite_principale", ""))
            break
        except requests.exceptions.RequestException:
            time.sleep(1)

    # Fallback with city
    if city and str(city).strip() not in ("", "None"):
        try:
            resp2 = session.get(API_URL, params={"q": f"{query} {city}", "per_page": 5}, timeout=15)
            if resp2.status_code == 200:
                results2 = resp2.json().get("results", [])
                if results2:
                    best = results2[0]
                    siege = best.get("siege", {})
                    return (best.get("siren", ""), siege.get("siret", ""),
                            best.get("nom_complet", ""), best.get("activite_principale", ""))
        except requests.exceptions.RequestException:
            pass

    return "", "", "", ""


def process_garage(g, total):
    name = g.get("name", "")
    postal = g.get("postal_code", "")
    city = g.get("city", "")
    cache_key = f"{name}|{postal}|{city}"

    with cache_lock:
        if cache_key in cache:
            siren, siret, matched_name, naf = cache[cache_key]
            with counter_lock:
                counter["cached"] += 1
                counter["done"] += 1
                if siren:
                    counter["found"] += 1
                else:
                    counter["not_found"] += 1
            g["siren"] = siren
            g["siret"] = siret
            g["nom_api"] = matched_name
            g["naf"] = naf
            return

    siren, siret, matched_name, naf = search_siret(name, postal, city)

    with cache_lock:
        cache[cache_key] = [siren, siret, matched_name, naf]

    with counter_lock:
        counter["done"] += 1
        if siren:
            counter["found"] += 1
        else:
            counter["not_found"] += 1
        done = counter["done"]

    if done % 200 == 0 or done == total:
        with cache_lock:
            with open(SIRET_CACHE_FILE, "w", encoding="utf-8") as f:
                json.dump(cache, f, ensure_ascii=False)
        print(f"  [{done:>6}/{total}] ({100*done/total:5.1f}%) | "
              f"Trouvé: {counter['found']} | Non trouvé: {counter['not_found']}", flush=True)

    g["siren"] = siren
    g["siret"] = siret
    g["nom_api"] = matched_name
    g["naf"] = naf


def main():
    global cache

    print("=" * 70)
    print("  EXPORT GARAGES PJ → EXCEL + SIRET (FAST)")
    print("=" * 70)

    garages = load_all_garages()
    total = len(garages)
    with_phone = sum(1 for g in garages if g.get("phone"))
    print(f"\nTotal: {total} garages, {with_phone} avec tel ({100*with_phone/total:.1f}%)")

    if os.path.exists(SIRET_CACHE_FILE):
        with open(SIRET_CACHE_FILE, "r", encoding="utf-8") as f:
            cache = json.load(f)
    print(f"Cache SIRET: {len(cache)} entrées")

    remaining = sum(1 for g in garages if f"{g.get('name','')}|{g.get('postal_code','')}|{g.get('city','')}" not in cache)
    print(f"Requêtes API restantes: {remaining}")
    print(f"Workers: {WORKERS}\n")

    with ThreadPoolExecutor(max_workers=WORKERS) as executor:
        futures = [executor.submit(process_garage, g, total) for g in garages]
        for future in as_completed(futures):
            try:
                future.result()
            except Exception as e:
                print(f"  [ERROR] {e}", flush=True)

    # Final cache save
    with open(SIRET_CACHE_FILE, "w", encoding="utf-8") as f:
        json.dump(cache, f, ensure_ascii=False)

    # Export Excel
    print(f"\nCréation Excel: {OUTPUT_FILE}...")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Garages Pages Jaunes"

    headers = [
        "N°", "Nom", "Téléphone", "Téléphone (Intl)", "Adresse",
        "Code Postal", "Ville", "Département N°", "Département",
        "Région", "Activité PJ", "SIREN", "SIRET", "Nom API",
        "Code NAF", "ID Pages Jaunes", "Source",
    ]
    for col, h in enumerate(headers, 1):
        ws.cell(row=1, column=col, value=h)

    for i, g in enumerate(garages):
        row = i + 2
        ws.cell(row=row, column=1, value=i + 1)
        ws.cell(row=row, column=2, value=g.get("name", ""))
        ws.cell(row=row, column=3, value=g.get("phone", ""))
        ws.cell(row=row, column=4, value=g.get("phone_intl", ""))
        ws.cell(row=row, column=5, value=g.get("address", ""))
        ws.cell(row=row, column=6, value=g.get("postal_code", ""))
        ws.cell(row=row, column=7, value=g.get("city", ""))
        ws.cell(row=row, column=8, value=g.get("dept_num", ""))
        ws.cell(row=row, column=9, value=g.get("dept_name", ""))
        ws.cell(row=row, column=10, value=g.get("region", ""))
        ws.cell(row=row, column=11, value=g.get("activity", ""))
        ws.cell(row=row, column=12, value=g.get("siren", ""))
        ws.cell(row=row, column=13, value=g.get("siret", ""))
        ws.cell(row=row, column=14, value=g.get("nom_api", ""))
        ws.cell(row=row, column=15, value=g.get("naf", ""))
        ws.cell(row=row, column=16, value=g.get("pj_id", ""))
        ws.cell(row=row, column=17, value="Pages Jaunes")

    wb.save(OUTPUT_FILE)

    found = counter["found"]
    not_found = counter["not_found"]
    print(f"\n{'=' * 70}")
    print(f"  Total: {total} | Avec tel: {with_phone} ({100*with_phone/total:.1f}%)")
    print(f"  SIRET trouvé: {found} ({100*found/total:.1f}%)")
    print(f"  SIRET non trouvé: {not_found}")
    print(f"  Fichier: {OUTPUT_FILE}")
    print(f"{'=' * 70}")


if __name__ == "__main__":
    main()
