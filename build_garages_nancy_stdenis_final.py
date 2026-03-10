#!/usr/bin/env python3
"""
Merge all garage data sources for Nancy (54) & Saint-Denis (93),
enrich missing SIRET via API Gouv, fix phone numbers, export final Excel.

Sources merged:
  - garages_nancy_saintdenis.json  (PagesJaunes scrape with SIRET)
  - scrape_nancy_pagesjaunes.json  (second PJ scrape)
  - scrape_nancy_annuaire.json     (annuaire scrape)
  - scrape_stdenis_annuaire.json   (annuaire scrape)
  - scrape_stdenis_societe.json    (societe.com legal data)
  - scrape_entreprises_legal.json  (legal enterprise data)
  - scrape_environs_legal.json     (API Gouv bulk legal data, depts 54+93)
"""

import json
import re
import time
import urllib.request
import urllib.parse
import ssl
import os

try:
    import openpyxl
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
except ImportError:
    print("ERROR: openpyxl required. Install with: pip install openpyxl")
    raise

BASE = "/home/user/Website01"
OUTPUT_EXCEL = os.path.join(BASE, "garages_nancy_saintdenis.xlsx")
OUTPUT_JSON = os.path.join(BASE, "garages_nancy_saintdenis_final.json")

# ── Phone number patterns (ordered by priority) ──
PHONE_PATTERNS = [
    re.compile(r'\+377\s\d{2}\s\d{2}\s\d{2}\s\d{2}'),
    re.compile(r'\+33\s[1-9](?:\s\d{2}){4}'),
    re.compile(r'0\s8\d{2}\s\d{2}\s\d{2}\s\d{2}'),
    re.compile(r'0[1-9](?:\s\d{2}){4}'),
    re.compile(r'0[1-9]\d{8}'),
]


def extract_phones(raw):
    """Extract all valid phone numbers from a possibly concatenated string."""
    if not raw:
        return []
    phones = []
    for pat in PHONE_PATTERNS:
        for m in pat.finditer(raw):
            phones.append(m.group())
    # Deduplicate preserving order
    seen = set()
    result = []
    for p in phones:
        normalized = re.sub(r'\s+', '', p)
        if normalized not in seen:
            seen.add(normalized)
            result.append(p)
    return result


def format_intl(phone):
    """Convert French local number to international format."""
    if not phone:
        return ""
    digits = re.sub(r'\s+', '', phone)
    if digits.startswith('+33'):
        return phone
    if digits.startswith('+377'):
        return phone
    if digits.startswith('0') and len(digits) == 10:
        intl = '+33 ' + digits[1]
        for i in range(2, 10, 2):
            intl += ' ' + digits[i:i+2]
        return intl
    return phone


def normalize_name(name):
    """Normalize name for dedup matching."""
    if not name:
        return ""
    n = name.upper().strip()
    n = re.sub(r'[^A-Z0-9 ]', ' ', n)
    n = re.sub(r'\s+', ' ', n).strip()
    return n


def make_key(entry):
    """Create dedup key from name + postal_code."""
    name = normalize_name(entry.get("name", ""))
    pc = (entry.get("postal_code", "") or "").strip()
    return f"{name}|{pc}"


def make_siren_key(entry):
    """Create key from SIREN if available."""
    siren = (entry.get("siren", "") or "").strip()
    return siren if siren else None


# ── SIRET enrichment via API Gouv ──
def fetch_url(url, max_retries=3):
    ctx = ssl.create_default_context()
    req = urllib.request.Request(url, headers={"User-Agent": "Mozilla/5.0"})
    for attempt in range(max_retries):
        try:
            with urllib.request.urlopen(req, context=ctx, timeout=30) as resp:
                return json.loads(resp.read().decode("utf-8"))
        except urllib.error.HTTPError as e:
            if e.code == 429:
                time.sleep(10 * (attempt + 1))
                continue
            return None
        except Exception:
            time.sleep(5)
            continue
    return None


def search_siret_api(name, postal_code, city=""):
    """Search for SIRET via recherche-entreprises.api.gouv.fr."""
    if not name:
        return None
    q = name
    params = {"q": q, "per_page": "5"}
    if postal_code:
        dept = postal_code[:2]
        params["departement"] = dept

    url = "https://recherche-entreprises.api.gouv.fr/search?" + urllib.parse.urlencode(params)
    data = fetch_url(url)
    if not data or "results" not in data or not data["results"]:
        return None

    norm_name = normalize_name(name)
    best = None
    best_score = 0

    for ent in data["results"][:5]:
        ent_name = normalize_name(ent.get("nom_complet", "") or ent.get("nom_raison_sociale", ""))
        siege = ent.get("siege", {})
        ent_cp = siege.get("code_postal", "") or ""

        # Score matching
        score = 0
        # Name similarity (simple word overlap)
        name_words = set(norm_name.split())
        ent_words = set(ent_name.split())
        if name_words and ent_words:
            overlap = len(name_words & ent_words)
            score = overlap / max(len(name_words), len(ent_words))

        # Bonus for postal code match
        if postal_code and ent_cp == postal_code:
            score += 0.3
        elif postal_code and ent_cp[:2] == postal_code[:2]:
            score += 0.1

        if score > best_score:
            best_score = score
            dirigeants = ent.get("dirigeants", [])
            director_parts = []
            for d in dirigeants:
                nom = d.get("nom", "")
                prenom = d.get("prenom", "")
                qualite = d.get("qualite", "")
                if nom:
                    dname = f"{prenom} {nom}".strip()
                    if qualite:
                        dname += f" ({qualite})"
                    director_parts.append(dname)

            etat = ent.get("etat_administratif", "")
            status = "Active" if etat == "A" else "Radiée" if etat == "C" else etat

            best = {
                "siren": ent.get("siren", ""),
                "siret": siege.get("siret", ""),
                "naf_code": ent.get("activite_principale", "") or siege.get("activite_principale", ""),
                "director": "; ".join(director_parts),
                "creation_date": ent.get("date_creation", ""),
                "status": status,
                "employees": ent.get("tranche_effectif_salarie", "") or "",
            }

    # Only accept if score is decent
    if best_score >= 0.4:
        return best
    return None


def load_json(filename):
    path = os.path.join(BASE, filename)
    if not os.path.exists(path):
        print(f"  WARNING: {filename} not found, skipping")
        return None
    with open(path, "r", encoding="utf-8") as f:
        return json.load(f)


def merge_entry(existing, new_entry):
    """Merge new_entry fields into existing, filling blanks."""
    for key in new_entry:
        new_val = new_entry.get(key, "") or ""
        old_val = existing.get(key, "") or ""
        if new_val and not old_val:
            existing[key] = new_val
    # Merge sources
    old_src = existing.get("source", "")
    new_src = new_entry.get("source", "")
    if new_src and new_src not in old_src:
        existing["source"] = f"{old_src}; {new_src}" if old_src else new_src


def dept_for(entry):
    """Get department from postal code."""
    pc = (entry.get("postal_code", "") or "").strip()
    return pc[:2] if len(pc) >= 2 else ""


def main():
    # Master dict: keyed by dedup key
    master = {}  # key -> entry
    siren_index = {}  # siren -> key

    def add_entry(entry, region_label=""):
        """Add or merge an entry into master."""
        key = make_key(entry)
        siren = make_siren_key(entry)

        # Check if we already have this by SIREN
        existing_key = siren_index.get(siren) if siren else None
        if existing_key and existing_key in master:
            merge_entry(master[existing_key], entry)
            return

        if key in master:
            merge_entry(master[key], entry)
            if siren:
                siren_index[siren] = key
        else:
            master[key] = dict(entry)
            if region_label:
                master[key].setdefault("region", region_label)
            if siren:
                siren_index[siren] = key

    # ── 1. Load main PagesJaunes scrape ──
    print("Loading garages_nancy_saintdenis.json...")
    pj_main = load_json("garages_nancy_saintdenis.json")
    if pj_main:
        for label, entries in pj_main.items():
            region = "Nancy" if "54" in label else "Saint-Denis"
            for e in entries:
                add_entry(e, region)
    print(f"  Master: {len(master)} entries")

    # ── 2. Load second PJ scrape (Nancy) ──
    print("Loading scrape_nancy_pagesjaunes.json...")
    pj2 = load_json("scrape_nancy_pagesjaunes.json")
    if pj2:
        for e in pj2:
            add_entry(e, "Nancy")
    print(f"  Master: {len(master)} entries")

    # ── 3. Load annuaire scrapes ──
    print("Loading annuaire scrapes...")
    for fname, region in [("scrape_nancy_annuaire.json", "Nancy"),
                           ("scrape_stdenis_annuaire.json", "Saint-Denis")]:
        data = load_json(fname)
        if data:
            for e in data:
                add_entry(e, region)
    print(f"  Master: {len(master)} entries")

    # ── 4. Load societe.com data ──
    print("Loading scrape_stdenis_societe.json...")
    soc = load_json("scrape_stdenis_societe.json")
    if soc:
        for e in soc:
            add_entry(e, "Saint-Denis")
    print(f"  Master: {len(master)} entries")

    # ── 5. Load legal enterprise data ──
    print("Loading scrape_entreprises_legal.json...")
    leg = load_json("scrape_entreprises_legal.json")
    if leg:
        for e in leg:
            dept = dept_for(e)
            region = "Nancy" if dept == "54" else "Saint-Denis" if dept == "93" else ""
            add_entry(e, region)
    print(f"  Master: {len(master)} entries")

    # ── 6. Load API Gouv bulk data (filter to depts 54, 93) ──
    print("Loading scrape_environs_legal.json (API Gouv bulk)...")
    api_bulk = load_json("scrape_environs_legal.json")
    if api_bulk:
        added = 0
        for e in api_bulk:
            dept = dept_for(e)
            if dept in ("54", "93"):
                region = "Nancy" if dept == "54" else "Saint-Denis"
                # Only add Active ones
                if e.get("status", "") == "Active":
                    add_entry(e, region)
                    added += 1
        print(f"  Added/merged {added} active enterprises from API Gouv bulk")
    print(f"  Master: {len(master)} entries")

    # ── 7. Enrich missing SIRET ──
    entries_list = list(master.values())
    missing_siret = [e for e in entries_list if not (e.get("siret") or e.get("siren"))]
    print(f"\nEnriching SIRET for {len(missing_siret)} entries without SIRET...")

    enriched = 0
    for i, e in enumerate(missing_siret):
        result = search_siret_api(e.get("name", ""), e.get("postal_code", ""), e.get("city", ""))
        if result:
            for k, v in result.items():
                if v and not e.get(k):
                    e[k] = v
            enriched += 1
        if (i + 1) % 20 == 0:
            print(f"  {i+1}/{len(missing_siret)} processed, {enriched} enriched")
        time.sleep(0.2)  # Rate limit

    print(f"  Enriched {enriched}/{len(missing_siret)} entries with SIRET")

    # ── 8. Fix phone numbers ──
    print("\nFixing phone numbers...")
    fixed_phones = 0
    for e in entries_list:
        raw_phone = e.get("phone", "") or ""
        # Detect concatenated phones (standard French phone is 14 chars with spaces)
        if len(raw_phone) > 16:
            phones = extract_phones(raw_phone)
            if phones:
                e["phone"] = phones[0]
                e["phone_intl"] = format_intl(phones[0])
                if len(phones) > 1:
                    e["other_phones"] = "; ".join(phones[1:])
                fixed_phones += 1
        # Ensure intl format exists
        if e.get("phone") and not e.get("phone_intl"):
            e["phone_intl"] = format_intl(e["phone"])

        # Also handle phone2 field
        phone2 = e.get("phone2", "") or ""
        if phone2 and not e.get("other_phones"):
            e["other_phones"] = phone2

    print(f"  Fixed {fixed_phones} concatenated phone numbers")

    # ── 9. Filter: only depts 54 and 93, only Active (or no status) ──
    final = []
    for e in entries_list:
        dept = dept_for(e)
        if dept not in ("54", "93"):
            continue
        status = e.get("status", "")
        if status and status == "Radiée":
            continue
        final.append(e)

    # Sort by department, city, name
    final.sort(key=lambda x: (
        dept_for(x),
        (x.get("city", "") or "").upper(),
        (x.get("name", "") or "").upper()
    ))

    print(f"\nFinal dataset: {len(final)} entries")
    dept54 = [e for e in final if dept_for(e) == "54"]
    dept93 = [e for e in final if dept_for(e) == "93"]
    with_phone = sum(1 for e in final if e.get("phone"))
    with_siret = sum(1 for e in final if e.get("siret"))
    with_both = sum(1 for e in final if e.get("phone") and e.get("siret"))
    print(f"  Dept 54 (Nancy area): {len(dept54)}")
    print(f"  Dept 93 (Saint-Denis area): {len(dept93)}")
    print(f"  With phone: {with_phone}")
    print(f"  With SIRET: {with_siret}")
    print(f"  With both: {with_both}")

    # ── 10. Save JSON ──
    with open(OUTPUT_JSON, "w", encoding="utf-8") as f:
        json.dump(final, f, ensure_ascii=False, indent=2)
    print(f"\nSaved JSON: {OUTPUT_JSON}")

    # ── 11. Export Excel ──
    print("Exporting Excel...")
    export_excel(final, dept54, dept93, with_phone, with_siret, with_both)
    print(f"Saved Excel: {OUTPUT_EXCEL}")
    print("\n=== DONE ===")


def export_excel(final, dept54, dept93, with_phone, with_siret, with_both):
    """Export to styled Excel workbook."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Garages"

    # Styles
    header_font = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill(start_color="1A5276", end_color="1A5276", fill_type="solid")
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

    alt_fill_54 = PatternFill(start_color="E8F8F5", end_color="E8F8F5", fill_type="solid")
    alt_fill_93 = PatternFill(start_color="EBF5FB", end_color="EBF5FB", fill_type="solid")
    white_fill = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")

    thin_border = Border(
        left=Side(style="thin", color="D5D8DC"),
        right=Side(style="thin", color="D5D8DC"),
        top=Side(style="thin", color="D5D8DC"),
        bottom=Side(style="thin", color="D5D8DC"),
    )

    # Headers
    headers = [
        ("N°", 6),
        ("Nom de l'établissement", 40),
        ("Adresse", 35),
        ("Code Postal", 12),
        ("Ville", 22),
        ("Département", 14),
        ("Téléphone", 20),
        ("Téléphone (International)", 26),
        ("Autres téléphones", 22),
        ("SIREN", 14),
        ("SIRET", 18),
        ("Code NAF", 10),
        ("Dirigeant", 30),
        ("Date de création", 16),
        ("Statut", 10),
        ("Effectifs", 10),
        ("Source", 25),
    ]

    for col_idx, (title, width) in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx, value=title)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        cell.border = thin_border
        ws.column_dimensions[openpyxl.utils.get_column_letter(col_idx)].width = width

    # Data rows
    dept_names = {"54": "Meurthe-et-Moselle (54)", "93": "Seine-Saint-Denis (93)"}

    for row_idx, entry in enumerate(final, 2):
        dept = dept_for(entry)
        values = [
            row_idx - 1,
            entry.get("name", ""),
            entry.get("address", ""),
            entry.get("postal_code", ""),
            entry.get("city", ""),
            dept_names.get(dept, dept),
            entry.get("phone", ""),
            entry.get("phone_intl", ""),
            entry.get("other_phones", ""),
            entry.get("siren", ""),
            entry.get("siret", ""),
            entry.get("naf_code", ""),
            entry.get("director", ""),
            entry.get("creation_date", ""),
            entry.get("status", ""),
            entry.get("employees", ""),
            entry.get("source", ""),
        ]

        fill = alt_fill_54 if dept == "54" else alt_fill_93 if dept == "93" else white_fill
        if row_idx % 2 == 0:
            fill = white_fill

        for col_idx, val in enumerate(values, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.border = thin_border
            cell.fill = fill
            if col_idx == 1:
                cell.alignment = Alignment(horizontal="center")

    # Freeze header row
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{openpyxl.utils.get_column_letter(len(headers))}1"

    # ── Summary sheet ──
    ws2 = wb.create_sheet("Résumé")
    ws2.column_dimensions["A"].width = 35
    ws2.column_dimensions["B"].width = 15

    summary_data = [
        ("RÉSUMÉ - Garages Nancy & Saint-Denis", ""),
        ("", ""),
        ("Total garages", len(final)),
        ("Dept 54 - Meurthe-et-Moselle", len(dept54)),
        ("Dept 93 - Seine-Saint-Denis", len(dept93)),
        ("", ""),
        ("Avec téléphone", with_phone),
        ("Avec SIRET", with_siret),
        ("Avec téléphone ET SIRET", with_both),
        ("", ""),
        ("Taux téléphone", f"{100*with_phone/len(final):.1f}%" if final else "N/A"),
        ("Taux SIRET", f"{100*with_siret/len(final):.1f}%" if final else "N/A"),
    ]

    title_font = Font(name="Calibri", bold=True, size=14, color="1A5276")
    label_font = Font(name="Calibri", bold=True, size=11)
    value_font = Font(name="Calibri", size=11)

    for row_idx, (label, value) in enumerate(summary_data, 1):
        cell_a = ws2.cell(row=row_idx, column=1, value=label)
        cell_b = ws2.cell(row=row_idx, column=2, value=value)
        if row_idx == 1:
            cell_a.font = title_font
        else:
            cell_a.font = label_font
            cell_b.font = value_font
            cell_b.alignment = Alignment(horizontal="center")

    wb.save(OUTPUT_EXCEL)


if __name__ == "__main__":
    main()
